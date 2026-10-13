from pathlib import Path
import json,csv,re,hashlib,zipfile
from docx import Document
from lxml import etree
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[3];OUT=Path(__file__).resolve().parent;BASE=OUT.parent;E=OUT/'evidence'
def write(n,v):(E/n).write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
manifest=json.loads((E/'source_manifest.json').read_text());by={s['path']:s['evidence_id'] for s in manifest['sources']}
refs=json.loads((BASE/'evidence/references.json').read_text())
for r in refs:r['number']+=1
refs.insert(0,{'number':1,'authors':'PIARC','title':'Road Tunnels Manual: Video Automatic Incident Detection','publication_type':'technical_manual','year':'','opened_primary_url':'https://tunnelsmanual.piarc.org/en/equipment-systems-monitoringincident-detection/automatic-incident-detection','support_locator':'opening paragraph; section 2 Design and commissioning','assistant_checked_on':'2026-10-03','human_verification':'pending'})
write('references.json',refs)
known={s['evidence_id'] for s in manifest['sources']}
rows=list(csv.DictReader((E/'claims.csv').open(encoding='utf-8-sig')));claim_by={r['claim_id']:r for r in rows}
audit=[]
for line in (E/'manuscript.md').read_text().splitlines():
    hit=re.search(r'\[claim:(C\d+)\]',line)
    if not hit:continue
    content=re.sub(r'\[(?:claim|evidence):[^\]]*\]','',line)
    h=hashlib.sha256(' '.join(content.split()).encode()).hexdigest()
    row=claim_by[hit[1]];assert row['claim_text_sha256']==h,hit[1]
    ids=[x for x in row['evidence_ids'].split(',') if x]
    assert all(x in known for x in ids),ids
    audit.append({'claim_id':hit[1],'hash_binding':'PASS','evidence_id_binding':'PASS','human_verification':'pending','claim_role':'interpretation_or_proposed_work' if any(x in content for x in ['后续可','后续应','可进一步','可以让','宜按','教学中可以','进一步研究']) else 'descriptive_or_method'})
write('claim-binding-audit.json',audit)
cm=json.loads((BASE/'evidence/consistency_manifest.json').read_text())
for n in cm['numeric_facts']:
    if n['section']=='4.4':n['section']='4.3'
for r in cm['results']:r['reported_sections']=['4.3','Table 5','Figure 7']
write('consistency_manifest.json',cm)
write('reporting_coverage.json',{'document_type':'Chinese undergraduate engineering competition project manual','structure':['design background and research objectives','system solution','design contributions','research practice and verification','application value and phased extension','conclusions'],'clinical_guideline':'not_applicable; engineering project description and descriptive software verification','topics':{'need_and_objectives':'1','source_models':'2.2 and Table 2','method_summary':'2.3–2.7','contribution_evidence':'3 and Table 3','undergraduate_research_process':'4.1 and Table 4','software_validation':'4.2','descriptive_outputs':'4.3, Table 5 and Figure 7','independent_real_truth':'missing, disclosed in 4.4','next_experiments':'4.5','application_conditions':'5 and Table 6','human_approval':'pending','team_members_and_individual_contributions':'pending'},'submission_ready':False})
displays=json.loads((E/'displays.json').read_text())
fs={1:['backend/processor.py','backend/service.py','frontend/src/App.tsx'],2:['README.md','docs/features-analysis.md','docs/visual-evaluation.md'],3:['backend/calibration/homography.py'],4:['output/playwright/twin-horizontal.png','docs/twin-orientation.md'],5:['backend/evidence.py','docs/features-analysis.md'],6:['plans/features-report-desktop.jpg'],7:['backend/data/runs/20261002-170236-c53990/traffic_statistics.csv'],8:['plans/002-evaluation-result.jpg','plans/002-browser-result.json']}
for d in displays:
    if d['kind']=='figure':
        n=d['number'];d['evidence']=[by[x] for x in fs[n]];d['asset_sha256']=sha(OUT/'assets'/d['asset'])
        d['processing']={1:'Retained source-derived architecture diagram',2:'New source-derived user and research workflow diagram; no experimental results generated',3:'Retained conceptual diagram; not a real scene calibration',4:'Retained crop of actual screenshot with original ID 1 insets enlarged; no factual alteration',5:'Retained source-derived evidence workflow',6:'Retained actual report screenshot crop',7:'Retained plot of all 256 actual CSV rows; no smoothing; no independent truth',8:'Retained controlled-test screenshot with explicit constructed-label disclaimer'}[n]
        d['rights_status']='User-provided project materials; competition redistribution rights remain for team confirmation'
write('displays.json',displays)
p=OUT/'基于视频感知的隧道交通数字孪生系统作品说明书.docx';doc=Document(p)
with zipfile.ZipFile(p) as z:
    xml=etree.fromstring(z.read('word/document.xml'));settings=etree.fromstring(z.read('word/settings.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
assert len(doc.tables)==6 and len(doc.inline_shapes)==8
assert not xml.xpath('//m:oMath',namespaces=ns)
assert not settings.xpath('//w:documentProtection',namespaces=ns)
cites=xml.xpath('//w:r[w:rPr/w:vertAlign[@w:val="superscript"]]/w:t/text()',namespaces=ns)
first=[]
for s in cites:
    v=re.search(r'\[(\d+)(?:-(\d+))?\]',s);assert v,s
    for n in range(int(v[1]),int(v[2] or v[1])+1):
        if n not in first:first.append(n)
assert first==list(range(1,11)),first
figs=[int(re.match(r'图(\d+)',q.text)[1]) for q in doc.paragraphs if re.match(r'图\d+  ',q.text)]
tables=[int(re.match(r'表(\d+)',q.text)[1]) for q in doc.paragraphs if re.match(r'表\d+  ',q.text)]
assert figs==list(range(1,9)) and tables==list(range(1,7))
drift=[]
for s in manifest['sources']:
    if not s.get('sha256'):continue
    path=Path(s['path']);path=path if path.is_absolute() else ROOT/path
    if not path.is_file() or sha(path)!=s['sha256']:drift.append(s['path'])
assert not drift,drift
summary=json.loads((E/'run-summary.json').read_text())
assert sha(ROOT/'backend/data/videos/demo.mp4')==summary['source_sha256']
assert summary['observed_rows']+summary['predicted_rows']==summary['vehicle_rows']==2162
assert summary['machine_crossings']==19 and summary['valid_observed_speed_samples']==0
pdf=next(Path('/private/tmp/twin-undergraduate-qa3').glob('*.pdf'));pages=len(PdfReader(pdf).pages)
assert pages<=20
write('delivery-audit.json',{'docx_sha256':sha(p),'pages':pages,'figures':8,'native_editable_tables':6,'editable_equations':0,'algorithm_equations_removed':True,'reference_count':10,'genuine_journal_papers':3,'superscript_citation_runs':len(cites),'references_ordered_by_first_use':True,'source_hash_mismatches':drift,'original_video_identity_rechecked':True,'claim_hash_and_evidence_id_binding':'PASS','numeric_consistency':'PASS','layout_review':'pending inspection of changed final pages','render_directory':'/private/tmp/twin-undergraduate-qa3','human_submission_approval':'pending','file_size_bytes':p.stat().st_size})
write('revision-record.json',{'date':'2026-10-03','request':'Retain original evidence but emphasize undergraduate research, background, design contribution and application value','original_docx_sha256':sha(BASE/'基于视频感知的隧道交通数字孪生系统作品说明书.docx'),'changes':['Rewrite abstract and background around need and research questions','Replace algorithm derivations with concise function principles','Expand design contributions with problem, implementation and evidence','Add undergraduate research process and inspectable outputs','Add concrete teaching tasks and phased tunnel application plan','Rebuild native tables around needs, contribution and research practice','Replace algorithm-heavy flow with user and research workflow','Retain real interface, controlled-test disclaimer and actual CSV curve','Retain genuine papers and superscript citations; add PIARC primary source'],'experiments_added':False,'source_data_modified':False,'team_contributions_inferred':False,'human_verification':'pending'})
check='''修订说明与证据来源及待补材料
2026年10月3日

一、本次修改
在原说明书基础上重写摘要、背景、研究目标、创新、科研实践和应用前景。算法公式已删除，技术部分保留必要原理与使用前提。新增需求与功能对应、设计特色及依据、科研过程成果和阶段推广条件；流程图改为使用与研究验证流程。正文与6张表均为Word可编辑内容，保留8幅图。旧稿完整保留。

二、证据来源
1. 原项目源码、models.json/config.yaml，以及docs/功能、精度、视觉评测、模板和朝向文档；plans/验收、tests/与frontend/tests/。本次重新核对原证据哈希，无漂移；未运行项目或新增实验，未修改源码、视频、模型、标定与历史数据。
2. 软件结果沿用2026年10月3日隔离记录：后端88项、前端27项通过，HTTP/WebSocket/正常退出检查通过。它们不是感知准确率。
3. 真实桥梁运行20261002-170236-c53990：0—8.5秒、256个统计帧、19次机器过线；1074观测＋1088预测＝2162车辆帧记录。未实测标定，有效观测速样本0。图7来自全部CSV，未经独立人工真值核验。
4. 图4为另一段桥梁录像的真实双视图与原ID局部放大；图6为实际报告；图8为受控人工评测验收，构造标签结果不属于真实交通精度。概念图与流程图均明确性质。
5. 主文10条真实来源，含IEEE Transactions on Intelligent Transportation Systems与Computer Vision and Image Understanding的3篇期刊论文，新增PIARC隧道手册原始来源。正文数字上标按首次出现顺序编号。来源与哈希、引用及图表记录在evidence/。

三、待补材料
团队、学校、指导教师及个人贡献说明；最终赛道和模板；合法真实隧道录像；机位与道路尺寸实测；完整独立过线和事件标签、参考测速；素材及模型使用授权；RTSP、长时、GPU、并发及现场验收。当前桥梁演示、二维展示和本机单会话的边界已保留。应用场景与后续实验均按已有能力或计划说明，没有虚构研究经历、准确率、经济效益或落地案例。
本版按用户20页内要求制作；若最终采用已查验的2026年官方附件2六页限制，需另做正式压缩版：https://itsnet.cn/wenjianxiazai/181.html 。团队仍需完成事实、贡献和正式提交的人工确认，证据记录不冒充人工签署。

四、编制工具与方法
Codex辅助本次撰写和排版，使用scientific-writing与Documents技能。编制方法引用保留在本清单，未混入交通技术文献：
Kassis T, Agarwal V, He Y, Patel D, Brueckner A M. Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents[EB/OL]. 2026. arXiv:2609.00065. DOI:10.48550/arXiv.2609.00065. https://arxiv.org/abs/2609.00065 。原始记录已重新打开核对；该项为方法工具预印本，不是本项目实验依据。
'''
(OUT/'修订说明与证据来源及待补材料清单.txt').write_text(check,encoding='utf-8')
print('PASS',pages,'pages; editable tables; superscript references 1–10; unchanged source hashes; original preserved.')
