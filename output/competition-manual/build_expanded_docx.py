from pathlib import Path
import csv,json,hashlib,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.5)
sec.left_margin=sec.right_margin=Cm(2)
sec.footer_distance=Cm(1.25)
for name,size in [('Normal',12),('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',12),('Caption',10.5)]:
    s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'STSong')
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        s._element.rPr.rFonts.attrib.pop(qn('w:'+key),None)
    if s._element.pPr is not None:
        for border in list(s._element.pPr.findall(qn('w:pBdr'))):s._element.pPr.remove(border)
    pf=s.paragraph_format; pf.line_spacing=1.15; pf.space_before=Pt(0); pf.space_after=Pt(4)
    if name.startswith('Heading'): pf.space_before=Pt(8); pf.space_after=Pt(4); s.font.bold=True
    if name=='Normal': pf.first_line_indent=Pt(24)
    else: pf.first_line_indent=Pt(0)
    if name=='Title':s.font.bold=True
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.core_properties.title='基于视频感知的隧道交通数字孪生系统作品说明书'
doc.core_properties.author=''; doc.core_properties.subject='智能交通竞赛作品说明书'
doc.core_properties.keywords='车辆检测 多目标跟踪 道路标定 二维数字孪生 人工评测'
doc.core_properties.comments=''
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
r=footer.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)

manifest=json.loads((OUT/'evidence/source_manifest.json').read_text())
E={x['path']:x['evidence_id'] for x in manifest['sources']}
claims=[]; draft=[]; displays=[]
def evidence(*paths): return [E[p] for p in paths]
def record(text,paths):
    ids=evidence(*paths) if paths else []
    c=f'C{len(claims)+1:03d}'
    claims.append({'claim_id':c,'claim_text_sha256':hashlib.sha256(' '.join(text.split()).encode()).hexdigest(),'evidence_ids':','.join(ids),'verification_status':'human_pending','assistant_checked_on':'2026-10-03'})
    draft.append(f'{text} [claim:{c}] [evidence:{",".join(ids)}]')
def add_cited_text(pa,text):
    for part in re.split(r'(\[\d+(?:[-,]\d+)*\])',text):
        run=pa.add_run(('\u2060'+part) if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',part) else part)
        if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',part):
            run.font.superscript=True; run.font.size=Pt(9)
def p(text,*paths,small=False):
    pa=doc.add_paragraph()
    add_cited_text(pa,text)
    if small:
        pa.paragraph_format.first_line_indent=Pt(0)
        for r in pa.runs:r.font.size=Pt(10.5)
    record(text,paths)
    if text.startswith('H为3×3'):pa.paragraph_format.keep_with_next=True
    return pa
def h(text,level=2):
    doc.add_paragraph(text,'Heading '+str(level));draft.append('## '+text)
def newpage():pass
def fig(n,name,caption,height):
    im=OUT/'assets'/name
    from PIL import Image
    w,hh=Image.open(im).size; width=min(17,height*w/hh)
    pa=doc.add_paragraph(); pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
    pa.alignment=WD_ALIGN_PARAGRAPH.CENTER; pa.paragraph_format.keep_with_next=True
    pic=pa.add_run().add_picture(str(im),width=Cm(width)); pic._inline.docPr.set('descr',caption)
    cp=doc.add_paragraph(f'图{n}  {caption}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    displays.append({'number':n,'kind':'figure','asset':name,'caption':caption,'processing':'local diagram or screenshot crop; no fact alteration'})
def table(n,title,headers,rows,widths,paths):
    cp=doc.add_paragraph(f'表{n}  {title}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for col,w in zip(tb.columns,widths):col.width=Cm(w)
    for cell,w in zip(tb.rows[0].cells,widths):cell.width=Cm(w)
    for cell,text in zip(tb.rows[0].cells,headers):cell.text=text
    for row in rows:
        cells=tb.add_row().cells
        for cell,w,text in zip(cells,widths,row):cell.width=Cm(w);cell.text=str(text)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF3');tcpr.append(sh)
            for pa in cell.paragraphs:
                pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0);pa.paragraph_format.line_spacing=1
                if ci==0:pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if n==1 and ri<len(tb.rows)-1:pa.paragraph_format.keep_with_next=True
                raw=pa.text
                pa.clear();add_cited_text(pa,raw)
                for r in pa.runs:
                    if not r.font.superscript:r.font.size=Pt(10.5)
                    r.bold=ri==0
    record('；'.join([title]+[' / '.join(str(x) for x in row) for row in [headers]+rows]),paths)
    displays.append({'number':n,'kind':'table','title':title,'native_word_table':True,'evidence':evidence(*paths)})
def eq(prefix,numerator,denominator,suffix):
    pa=doc.add_paragraph();pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0)
    m=OxmlElement('m:oMath')
    def mr(t):
        r=OxmlElement('m:r'); x=OxmlElement('m:t');x.text=t;r.append(x);return r
    m.append(mr(prefix));f=OxmlElement('m:f')
    for tag,t in [('num',numerator),('den',denominator)]:
        z=OxmlElement('m:'+tag);z.append(mr(t));f.append(z)
    m.append(f);m.append(mr(suffix));pa._p.append(m)

# Page 1: front matter and scope, matching the official no-cover requirement.
pa=doc.add_paragraph('基于视频感知的隧道交通数字孪生系统','Title');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['参赛团队：待补充    指导教师：待补充','学校及院系：待补充    所在城市及邮编：待补充']:
    pa=doc.add_paragraph(text);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
pa=doc.add_paragraph('作品内容简介','Heading 1');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('本作品面向隧道交通监测中视频信息难以转化为可解释状态、异常报警缺少复核依据等问题，构建本机视频感知与二维道路数字孪生系统。系统采用现成预训练车辆检测模型，以SORT风格的Kalman预测和Hungarian匹配维持轨迹ID；用轨迹框底边中心定位车辆，在可靠实测尺度下通过四点Homography映射至米制平面，并利用实际视频时间戳和窗口最小二乘拟合计算速度。系统实现有限计数线过线统计、车道和类别分析、10/30/60秒车道窗口，以及停车、急减速和逆行规则。FastAPI与WebSocket发布同帧图像和状态，React展示视频、位置、朝向、轨迹、统计与事件；统计段归档、报警截图、独立回放、机位模板和版本化人工评测支持结果追溯。已有桥梁俯拍素材的运行记录与软件回归支撑功能实现，尚无真实隧道真值及实测测速精度结果。作品未训练或改进检测网络，当前为二维展示和本机单分析会话。未实测标定时，米制位置、速度、加速度与相关报警保持不可用。','backend/processor.py','backend/service.py','backend/vision/detector.py','docs/features-analysis.md','docs/visual-evaluation.md','plans/002-verification.md')
p('关键词：视频车辆检测；多目标跟踪；道路标定；二维数字孪生；证据复核',small=True)
h('1 研究背景',1)
h('1.1 隧道交通监测需求与应用问题')
p('隧道监测需将视频中的车辆在场、过线和疑似异常联系到道路空间。缺少尺度、时间和原画面，测速与报警难以解释；仅复核已报事件也无法发现完全漏报，因此系统同时提供独立全区间核验。','docs/features-analysis.md','docs/accuracy-evaluation.md')
p('隧道空间相对封闭，车辆进入、停留与行驶方向需要结合明确道路位置判断。本作品选择视频感知作为入口，目标是把既有画面组织为可查看的车辆状态、断面统计和疑似事件证据。面对运营人员或竞赛评委，系统应回答车辆位于哪条车道、统计针对哪个时间段、报警由哪项规则产生，以及能否返回原视频核验，而不是只输出一个未经解释的数值。','docs/features-analysis.md','README.md')
p('从工程实现看，检测框属于图像坐标，跟踪ID属于一次分析中的关联结果，米制速度则需要道路尺度与有效时间。三者不能相互替代。同样，检测到的轨迹数、穿越计数线的次数与人工观察到的物理车辆总数也可能不同。作品把这些口径显式保存在状态和导出数据中，使使用者能够识别结果的含义与适用范围。','backend/processor.py','backend/exporter.py','docs/accuracy-evaluation.md')
h('1.2 相关研究与技术现状')
p('交通深度学习综述强调时空及场景差异[1]；单目测速须可靠的场景尺度[2]和独立速度真值[3]。本作品采用实测四点标定，未采用文献[2]的三维模型自动尺度推断。YOLO11和SSDLite提供预训练检测组件[4-6]；VisDrone研究聚焦无人机检测与跟踪[7]，为俯拍选型提供依据。SORT结合Kalman滤波与Hungarian关联[8]。本作品采用已有方法，重点实现测量门控、同帧展示和可追溯评测。','models.json','backend/vision/detector.py','backend/vision/tracker.py')
p('相关研究给出的测速结果建立在其自身标定、数据和独立参考测量上，不能直接移用到本项目。Sochor等建立了具有独立速度测量的单目测速数据集[3]，说明“能够显示速度”与“速度误差已经验证”是不同阶段。本作品目前完成前者所需的软件路径和未标定保护，但演示桥梁没有可靠实测尺度，因而不报告真实测速误差。','docs/accuracy-evaluation.md','backend/traffic/speed_estimator.py')
p('预训练模型的发布资料用于确认权重来源与可选网络，论文用于解释已有技术和评测原则。发布者在其数据集上的性能不能作为作品在桥梁或隧道上的准确率；采用YOLO11n也不表示本作品训练了YOLO11网络。本文只将与源码相符的组件、流程和已有证据写为已实现内容。','models.json','backend/vision/detector.py','docs/accuracy-evaluation.md')
h('1.3 本作品的目标及实现范围')
p('当前完成本机录像分析与复核，主要素材为桥梁俯拍，隧道场景待验证。孪生为二维状态映射。底层读取保留RTSP接口，现有入口以文件为主，实时流未验收。','backend/vision/video_reader.py','backend/service.py','plans/002-verification.md')
p('实现范围包括车辆检测与ID跟踪、人工道路几何配置、满足前提时的定位测速、流量和车道统计、二维同步展示、异常规则、历史报告、证据复核、机位模板及人工过线/事件评测。人工速度真值仍需独立设备或独立距离时间核验；三维场景、车辆动力学、跨镜头重识别、多摄像头并发与生产级权限管理不在当前完成范围内。','docs/features-analysis.md','docs/visual-evaluation.md','backend/service.py')
p('本文将四类证据分别使用：当前源码说明功能如何实现；软件测试检查契约和边界；保存的桥梁运行输出说明真实素材能够经过流程；人工真值评测才回答计数、测速或报警误差。计划文件中提出的功能若未找到对应实现与验收，不作为已完成工作；历史通过数量采用当时记录，本次回归数量另行标明。','plans/002-verification.md','plans/2026-10-02-implementation-status.md','docs/accuracy-evaluation.md')
h('2 设计原理与关键技术',1)
h('2.1 系统总体架构与数据流程')
p('图1为单分析会话的源码架构，表1列出模块分工。','backend/service.py','backend/processor.py','backend/main.py')
fig(1,'fig1-architecture.png','系统总体架构  依据当前源码绘制',3.5)
p('输入层保存视频内容身份、道路配置和模型参数。分析层依次解码、执行间隔检测、更新跟踪、计算接触点，再根据观测与标定状态决定是否进行测速、计数及规则判断。输出层一方面发布完整帧快照，另一方面写入统计段与事件事实。浏览器使用相同快照展示视频与孪生；历史报告和人工评测读取已经保存的记录。','backend/processor.py','backend/service.py','backend/exporter.py')
p('后端采用整帧处理与导出锁，使某一帧的车辆、统计和事件来自一致处理结果。分析会话与统计段是不同对象：会话承担当前播放和处理，统计段承担有固定来源与条件的记录。更换视频、定位或调整关键条件后形成新的统计段，避免把不同时间区间或不同标定条件拼接为一组指标。','backend/processor.py','backend/service.py','docs/features-analysis.md')
table(1,'系统模块及功能',['模块','主要功能','输出与约束'],[
['视频与感知','解码PTS、检测适配、ID跟踪','区分实际观测和预测'],
['道路与交通','四点标定、测速、过线及窗口统计','米制测量须可靠实测尺度'],
['展示与证据','双视图同步、事件截图及独立回放','二维、本机单分析会话'],
['报告与评测','CSV/HTML、模板、人工标注与版本评测','机器结果、复核与真值分别存储']],[2.8,7.1,7.1],['backend/processor.py','backend/reports.py','backend/evidence.py','backend/annotations.py'])

newpage()
h('2.2 视频车辆检测与多目标跟踪')
h('2.2.1 检测模型来源与输入适配',3)
p('预训练权重见表2，加载前按models.json校验SHA256。类别统一为car、motorcycle、bus、truck，VisDrone的van并入car。桥梁初期抽查SSDLite漏检较多，默认选用VisDrone YOLO11n；尚无统一真值的模型精度对比。','backend/vision/detector.py','backend/vision/weights.py','models.json','README.md')
table(2,'预训练模型来源与选型依据',['模型','来源与权重','场景与选型依据'],[
['VisDrone YOLO11n','dronefreak的VisDrone微调权重[4]','俯拍小目标演示默认；由发布者训练'],
['COCO YOLO11n','Ultralytics官方yolo11n.pt[5]','通用道路视角备选；需新素材复核'],
['SSDLite MobileNetV3','Torchvision官方COCO权重[6]','320×320轻量推理；当前俯拍抽查漏检']],[4,6.2,6.8],['models.json','backend/vision/detector.py','README.md'])
p('默认VisDrone YOLO11n来自dronefreak发布的微调权重，适合当前俯拍小目标演示的选择依据是项目已有素材抽查，不是本地重新训练。COCO YOLO11n来自Ultralytics官方权重，可作为通用交通视角的备选；SSDLite采用Torchvision提供的COCO权重和320×320输入。三种适配器统一输出位置、类别与置信度，便于沿用后续同一跟踪与分析流程。','models.json','backend/vision/detector.py')
p('模型选型同时考虑视角、目标像素尺寸和本机处理成本。当前默认配置置信度阈值0.40、YOLO输入尺寸960、CPU线程数4；分块检测选项存在，但默认tiled=false，本说明书不据此宣称已验证分块检测提升。降低阈值可能增加误检，提高阈值可能漏掉小目标；调整输入尺寸会改变计算量和可见细节，最终取舍仍需同一真值区间的对照。','config.yaml','backend/vision/detector.py')
p('实际解码画面按最长边不超过960像素缩放供处理与编码，分析后的框和接触点按横纵缩放比恢复到原图像素。道路配置、原视频叠加与导出的像素坐标以原始尺寸为基础，避免将处理分辨率误当原视频尺寸。权重由本地文件加载并校验摘要；这个校验能核对文件一致性，不能验证模型精度。','backend/vision/video_reader.py','backend/processor.py','backend/vision/weights.py')
h('2.2.2 Kalman预测与轨迹关联',3)
p('默认每隔2帧检测，其间预测（图2）。框中心、宽高及变化率组成8维状态，观测为4维；IoU门控后以1−IoU代价进行Hungarian一对一关联。满足最少命中数才输出，超龄清理。没有外观重识别，遮挡和ID分裂可能导致重复或漏计。','backend/vision/tracker.py','config.yaml')
fig(2,'fig2-pipeline.png','视频感知与交通分析流程  观测门控区分检测与预测',2.7)
p('observed=true为检测关联后的轨迹，false为预测；预测用于展示，不更新测速、计数或规则。跟踪保留30帧，展示另有限制漏检时长，避免长期预测框输出。','backend/processor.py','backend/vision/tracker.py')
p('跟踪状态由框中心cx、cy、宽w、高h及其四个变化量组成。预测阶段按逐帧常速度模型推进，观测阶段用匹配检测框修正。状态转移中的步长按一次处理帧计算，变化量属于框几何的逐帧变化，并非米每秒；实际物理速度另外由道路点与解码时间计算。对可变帧率素材，这一区分尤为必要，不能把Kalman内部变化量直接当作真实车速。','backend/vision/tracker.py','backend/traffic/speed_estimator.py')
p('IoU定义为两框交集面积与并集面积之比。当前关联先排除IoU低于0.20的候选，再进行Hungarian全局分配，防止明显不相交的候选占用本可成立的匹配。未分配检测创建新ID，未关联轨迹继续预测并增加漏检年龄；默认命中至少2次才进入输出，年龄超过30帧移除。','backend/vision/tracker.py','config.yaml')
p('轨迹内部保留与界面输出使用不同期限。处理器还要求漏检年龄不超过min(max_age,2×检测间隔)，默认即4帧，避免把长时间没有观测的框持续展示。轨迹类别采用检测置信度约束减少低置信度抖动，但未加入外观重识别，因此类别稳定或ID连续都不是物理车辆身份正确的证明。','backend/processor.py','backend/vision/tracker.py')
h('2.2.3 实际观测与预测位置的使用边界',3)
p('实际观测指本帧检测与轨迹成功关联后的滤波位置，仍会受到检测框误差影响；预测指未得到本帧观测时的状态外推。两者均可用于可视化连续性，但只有实际观测追加分析轨迹、更新过线、测速与事件规则。CSV保存observed字段，使报告能够剔除预测行，不能仅因某个预测行保留了最近速度，就将它视为新的独立测量样本。','backend/processor.py','backend/reports.py')
p('ID分裂可能使同一物理车辆形成两个确认轨迹；遮挡后重新出现也可能获得新ID。计数去重是“同一段同一轨迹ID一次”，不能保证“同一物理车辆一次”。因此28个确认ID不等于28辆真实车，98.9%的相邻帧ID保留率也不等于MOT准确率。物理身份正确性必须通过独立逐车或逐帧真值核验。','backend/vision/tracker.py','backend/traffic/flow_counter.py','docs/accuracy-evaluation.md')
h('2.3 道路标定与交通参数计算')
h('2.3.1 四点道路标定与米制坐标',3)
p('将轨迹框还原至原图，取底边中心近似道路接触点。依次选左上、右上、右下、左下角，并输入对应矩形的实测宽W、长L，求解H[9]。仅有效区域内的点参与米制测量；图3为概念示意。','backend/processor.py','backend/calibration/homography.py')
fig(3,'fig3-calibration.png','四点道路标定与坐标映射  概念示意，非实测道路',3.8)
eq('X = ','h₁₁u + h₁₂v + h₁₃','h₃₁u + h₃₂v + h₃₃','； Y = (h₂₁u + h₂₂v + h₂₃) / (h₃₁u + h₃₂v + h₃₃)')
p('四个图像点对应道路平面矩形的四个角，H将齐次像素坐标映射到道路坐标；尺度由W、L确定，不由检测模型推断。检测框底边中心近似轮胎与道路的接触位置，适合解释车辆在道路上的位置，但它并非真实接触点测量。在遮挡、俯视、检测框偏移或车辆姿态变化时，该近似会带来定位和测速误差。','backend/calibration/homography.py','backend/processor.py')
p('程序要求角点为有限数值、构成不交叉的凸四边形，面积与矩阵满足有效性条件，并检查道路尺寸范围。通过这些检查仅表示几何计算可进行，不表示角点对应正确、道路平面成立或尺寸已经实测。前端输入和模板复制都须由操作者确认尺度来源；本作品没有利用无人机高度、焦距或平均车长自动补出真实尺寸。','backend/calibration/homography.py','docs/camera-templates.md')
p('车辆点位于标定多边形内才输出米制位置；离开区域时不外推为有效速度。道路坡度、拱形路面、镜头畸变及相机运动会破坏单平面固定映射。应用到隧道前应保存机位、四角实物依据、测量方法与尺寸记录，并用未参与四点设置的独立测量位置检查映射偏差；这些属于待完成的实测验证。','backend/calibration/homography.py','backend/processor.py','docs/accuracy-evaluation.md')
h('2.3.2 实际时间戳与窗口最小二乘测速',3)
p('H为3×3单应矩阵，u、v以px计，X、Y以m计。对实际观测的米制点，令时间偏差Δtᵢ=tᵢ−平均时间，以窗口最小二乘拟合二维速度：','backend/calibration/homography.py','backend/traffic/speed_estimator.py')
eq('v = ','Σ Δtᵢ (rᵢ − r̄)','Σ (Δtᵢ)²','； s = 3.6 × mean₃(‖v‖)')
p('rᵢ为道路点，v单位m/s，s单位km/h。默认10个观测点，至少5点且跨度≥0.1s才测速，最近3次模长取均值；加速度由相邻平滑速度差计算。采用解码PTS，断档过长或离开有效区域时重置缓存。平面假设、畸变、相机移动和接触点偏差均待实测验证。','backend/traffic/speed_estimator.py','backend/processor.py','backend/vision/video_reader.py')
p('窗口最小二乘以时间为自变量、二维米制位置为因变量，利用多个观测点估计速度向量，减少逐点差分对框抖动的放大。smoothing_frames=10在实现中表示最多10个实际观测点，检测间隔改变时覆盖的秒数也随之变化，并非固定10个视频帧。最少5点与0.1s跨度用于避免样本不足和时间分母过小；最近3次速度模长再求均值会降低抖动，也引入一定响应延迟。','backend/traffic/speed_estimator.py','backend/processor.py')
p('视频时间来自解码器提供的PTS，要求有限、非负、严格单调。当前读取器遇到不可信时间戳时明确报错，而不以处理耗时或平均FPS合成时间。录制视频定位采用原始帧序号探测实际PTS，选择第一个不早于目标时间的帧，避免可变帧率输入按平均FPS定位到错误帧。原视频帧率仍用于元数据与部分近似步进，不承担真实测速的时间分母。','backend/vision/video_reader.py','tests/test_video_timing.py')
p('观测间隔超过max(0.5s,2.5×检测间隔/视频FPS)时重置测速及相关规则缓存，避免跨较长遮挡直接计算速度；离开有效标定区域也会清空测速历史。当前规则在重新累积足够样本前保持不可用。未标定记录中的空值表示测量前提不成立，不能改写为0km/h，因为0代表已测得静止，会错误触发停车判断。','backend/processor.py','backend/traffic/speed_estimator.py')
p('导出的加速度由相邻平滑速度差除以时间差计算，单位m/s²；急减速规则另用至少0.5s的速度变化以降低单帧波动。二者计算跨度不同，报警阈值不应机械地应用到CSV逐帧加速度列。速度和加速度是否准确仍依赖标定、接触点、时间与轨迹，软件公式正确不能代替独立参考测量。','backend/traffic/speed_estimator.py','backend/events/rules.py')
h('2.3.3 过线计数 车道归属与时间窗口',3)
p('过线判定检查连续实际观测点的连线与有限计数线相交，同一ID在一段内只计一次；车道以接触点落入人工多边形确定，类别和车道过线统计取穿越时归属。累计值不逐帧求和。10/30/60s窗口取(t−w,t]，覆盖为min(w,t−段起点)，不足时标记不完整；每分钟流率在不足60s时为短时折算。','backend/traffic/flow_counter.py','backend/processor.py','docs/features-analysis.md')
p('计数线是有限线段。程序检查同一ID相邻观测点之间的运动线段与它相交，排除仅穿越延长线的车辆以及平行移动。首次有效穿越时记录当前视频时间、车辆类别和车道，并将该ID加入去重集合。计数时间对应首次观测到跨线的采样时刻，没有按运动速度内插到精确穿越时刻；人工评测容差应考虑采样间隔和标注误差。','backend/traffic/flow_counter.py','backend/processor.py')
p('车道由人工多边形定义，接触点落入哪一多边形即归属该车道；未落入已定义车道的目标保留为“未划分”，不自动补画道路或推断车道。逐帧在场类别分布与过线类别累计分布是不同统计量：前者描述当前显示轨迹，后者描述已通过断面的ID。变道时的过线类别和车道取触发时结果，人工真值应采用同一统计口径。','backend/traffic/lane_detector.py','backend/traffic/flow_counter.py','backend/traffic/traffic_metrics.py')
p('10/30/60s车道窗口计数按当前时刻之前的指定区间聚合过线记录，并报告窗口长度、实际覆盖秒数与是否完整。伴随窗口行保存的观测数、预测数和平均速度是当前采样时刻状态，其中均速仅取实际观测且速度有效的车辆；它不是对整个10/30/60s时间段内所有车辆求得的窗口均速。阅读CSV时必须区分“过线的时间聚合”和“采样时刻的在场状态”。','backend/processor.py','backend/traffic/flow_counter.py')
p('每分钟流率将最近60s内过线数按实际覆盖时间折算；未覆盖完整60s时明确标记估计。累计流量取每帧累计列的最大值，若把它逐帧相加将重复计算同一过线。平均车头时距使用首末过线时间及穿越次数，只有超过一次穿越才可计算；这些指标均描述配置断面，不能直接解释为整段隧道的容量或交通服务水平。','backend/traffic/flow_counter.py','backend/reports.py')

newpage()
h('2.4 视频与二维孪生同步')
p('后端整帧完成后发布帧号、视频时间、JPEG、车辆、轨迹、统计和事件，WebSocket按修订号推送。前端以会话与修订号过滤旧响应，解码图像与来源快照成对提交，双视图共用状态，防止迟到图像错配位置。图4为未标定横向桥梁的实际同帧界面。','backend/service.py','backend/main.py','frontend/src/utils/snapshotState.ts','frontend/src/App.tsx')
p('状态传输以一帧为单位，不分别让前端独立拉取图像和车辆坐标。后端修订号表示发布状态的更新顺序，会话标识区别不同分析对象；前端抛弃旧会话、旧修订与晚到的解码结果。图像解码存在异步完成次序，若只替换坐标而图像仍停在旧帧，即会产生视觉误差。当前成对提交机制与乱序测试针对这个问题。','backend/main.py','frontend/src/utils/snapshotState.ts','frontend/tests/snapshotState.test.mjs')
p('同帧同步保证已提交图像与状态对应，并不保证浏览器显示每一原视频帧。推送、JPEG解码、渲染和本机负载可能影响刷新速度；处理吞吐、视频PTS和显示帧率是三个不同量。系统将检测耗时、跟踪耗时和处理FPS用于运行状态观察，当前没有完成浏览器显示帧率或端到端时延的独立测量。','backend/processor.py','backend/main.py','frontend/src/App.tsx')
fig(4,'fig4-twin.png','桥梁视频与二维孪生同帧对照  实际界面及同一ID局部放大，未标定',9.6)
p('未标定时保留原图比例，标定后以米尺度等比例展示，按角点调整方向和坐标轴。朝向由近期观测位移估计，静止时不等于真实车头方向。实时全局均速可包含预测轨迹保留的最近速度；报告与车道窗口只用实际观测，二者口径不同。','frontend/src/utils/twinGeometry.ts','frontend/src/components/DigitalTwin.tsx','backend/processor.py','backend/traffic/traffic_metrics.py','backend/reports.py')
p('二维孪生采用统一坐标映射绘制道路、车道、车辆和轨迹。未配置道路时仅显示图像范围，避免用默认车道制造场景信息；标定后使用道路米制坐标并保持比例。横向、竖向和旋转道路的朝向变换与位置变换共用几何依据，实际截图可用于检查对应关系。它是对观测状态的二维映射，不是道路三维重建、交通行为模拟或完整物理仿真。','frontend/src/utils/twinGeometry.ts','frontend/src/components/DigitalTwin.tsx','docs/twin-orientation.md')
p('图4来自3840×2160横向桥梁素材的实际生产界面，与图8使用的2160×3840竖向记录不是同一个运行。横向图用于方向、比例及同帧对应的功能展示，竖向CSV用于统计口径说明；本文不把不同素材的截图和数据组合为同一组实验。横向车道为人工描绘的图像轮廓，没有加入可靠实测尺度。','docs/twin-orientation.md','output/playwright/twin-horizontal.png')
h('2.5 异常事件与证据复核')
p('异常停车、急减速和逆行阈值见表3，仅处理实际观测；无有效速度时清缓存并返回无报警，均依赖标定和足够观测。停车不含停车区语义，急减速受测速抖动影响；逆行需正确车道方向，ID变化可能中断判断。','backend/events/rules.py','backend/processor.py')
p('停车判断保存该ID开始低于3km/h的时刻，低速累计达到3s才产生候选，速度回升则清除计时。规则能发现“持续低速”的目标，但不能区分拥堵排队、许可停车、紧急停车或正常等待。隧道落地需要结合停车区域、运行工况及人工复核；当前没有实现停车区语义或事故因果诊断。','backend/events/rules.py')
p('急减速缓存速度样本，用至少0.5s跨度的速度差计算加速度，低于−4m/s²产生候选。采用较长时间差是为抑制单帧峰值，仍不能消除错误标定、遮挡后框跳动和检测接触点偏差。该规则仅识别速度变化模式，不等于碰撞预测或危险行为责任判定。','backend/events/rules.py')
p('逆行将短时世界轨迹位移与车道合法方向比较，方向余弦小于−0.6且速度超过3km/h时开始计时，持续至少1s才触发。默认方向可由各车道覆盖，未归属车道时使用全局方向。低速方向不稳定、车道设置错误或ID中断均可能影响判定。合法方向必须与当前标定和道路视图一致，不能从屏幕左右方向直接推断。','backend/events/rules.py','backend/processor.py','config.yaml')
table(3,'关键参数单位与启用条件',['参数','默认值','条件或统计含义'],[
['检测与匹配','置信度0.40；间隔2帧；IoU 0.20','输入960；最少命中2次'],
['测速窗口','10点；至少5点；跨度≥0.1s','实测四点尺度；输出km/h、m/s²'],
['异常停车','速度<3km/h，持续≥3s','有效速度与实际观测'],
['急减速','加速度<−4m/s²','规则使用≥0.5s速度差'],
['逆行','余弦<−0.6，持续≥1s','速度>3km/h；合法方向已配置'],
['冷却与车道窗口','同ID同事件5s；10/30/60s','视频时间；窗口不足标记不完整']],[3.2,6.3,7.5],['config.yaml','backend/events/rules.py','backend/traffic/speed_estimator.py','backend/traffic/flow_counter.py'])
p('同一ID同一事件类型默认冷却5s，满足持续条件的目标不会在每帧形成新事件，但冷却结束后仍可能再次报警。因此事件CSV的条数是规则触发次数，不是事故数，也不是异常车辆数。画面中的短时警示状态与持久保存的事件记录分别服务实时浏览和后续复核，历史事件不会因警示标记消失而删除。','backend/events/rules.py','backend/processor.py')
p('如图5，报警触发时原子保存稳定事件ID、观测目标截图和不可变事实，记录视频身份、时间、规则、轨迹与标定。独立播放器先核验原视频内容身份，再通过Range回放事件邻域；来源缺失或改变时禁止回放，截图仍保留。人工确认、误报及备注另存，不修改机器事件或实时会话。','backend/evidence.py','docs/features-analysis.md')
fig(5,'fig5-evidence.png','异常事件证据与人工复核流程  依据源码绘制',2.1)
p('报警证据以事件发生时的事实为基础保存，关联统计段、事件ID、帧号、时间、车辆与规则条件。机器事件CSV、截图与事实元数据用于重现“系统当时看到了什么”；人工确认、误报判断和备注另存，回答“复核者如何解释它”。后续修改复核状态不改变机器规则曾经触发这一事实。','backend/evidence.py','backend/exporter.py')
p('独立回放先按统计段保存的视频ID和内容身份核验原文件，然后通过HTTP Range读取邻域。该播放器与当前分析播放器分开，查看历史事件不定位或暂停实时分析。若原视频丢失或内容改变，禁止把同名替换文件当作原来源继续回放；已保存截图和复核记录仍可查看，但不能据此作新的完整区间真值声明。','backend/evidence.py','backend/annotations.py','docs/visual-evaluation.md')
p('确认已报事件只能检查所见报警，发现不了系统完全没有报出的真实事件。人工评测因此提供完整录像独立查看、过线记录和未映射事件入口。报警复核与科学评测共享来源追溯思路，但数据对象和完整性声明不同，不能直接把“已确认报警数”当作TP或把未确认报警全部当作FP。','docs/visual-evaluation.md','backend/evaluation.py')

newpage()
h('2.6 历史报告 机位模板和人工评测')
h('2.6.1 历史统计段与分析报告',3)
p('统计段保存来源、参数、道路与标定快照。暂停继续保留原段，定位或换条件形成新段，旧CSV保留。报告流式读取CSV，累计过线取最大值；均速以有效观测车辆帧为样本，覆盖率为有效测速样本数/实际观测车辆帧数，不是每车均速或检测准确率。图6为桥梁报告区间与累计值，未标定速度留空。','backend/service.py','backend/exporter.py','backend/reports.py','docs/features-analysis.md')
fig(6,'fig6-report.png','桥梁报告指标区  实际界面，19次机器过线',3.0)
p('CSV及离线HTML/SVG可导出。对照要求同视频身份、区间差≤0.1s、几何与尺度相同，另列模型和规则差异；数量变化不表示精度提升。趋势图降采样，原CSV保留明细。','backend/reports.py','docs/features-analysis.md')
p('运行目录保存run.json和按用途划分的CSV。run.json记录来源身份、模型与分析配置、道路几何、标定状态、统计段起止及完成状态；traffic_statistics.csv保存逐帧整体统计，vehicles.csv保存逐车辆帧明细，events.csv保存机器事件，lane_statistics.csv保存车道采样与窗口。原始记录比趋势图更适合核查单位、样本和条件。','backend/exporter.py','backend/processor.py')
p('报告速度平均采用所有有效观测车辆帧，不是先对每车平均再对车辆平均，因此长时间在画面中的车辆会贡献更多样本。覆盖率反映观测样本中有多少满足测速条件；有效样本少时即使存在速度均值，也不能认为整段交通速度已经充分测量。条件对照必须保留不同参数，模型输出数量增加既可能来自漏检减少，也可能来自误检或ID分裂增加。','backend/reports.py','docs/features-analysis.md')
h('2.6.2 固定机位道路模板与尺度复用',3)
p('固定机位模板保存已保存的几何和可选实测尺度，按版本管理。同比例分辨率按像素宽高比映射，宽高比偏差超过1%仅允许预览；默认应用几何会清除真实尺度。即使图像比例相同，也必须人工确认机位、视野和实测尺寸仍有效，才能单独复制尺度。裁切、旋转、移动或变焦后须重新核验，模板不能代替标定。','backend/camera_templates.py','docs/camera-templates.md')
p('模板解决相同固定机位重复描绘车道和计数线的问题。只有已经保存的道路几何进入模板，当前未提交草稿不混入模板事实；名称、备注和修订版本用于识别与冲突检查。同比例分辨率转换按坐标比例映射，在独立预览中确认后应用。默认只复制几何并清除原真实尺度，减少把另一机位测量条件无意带入当前视频的风险。','backend/camera_templates.py','docs/camera-templates.md')
p('受控验收检查了200×400到100×200的坐标减半、计数线y120到y60、超过1%的宽高比偏差拒绝应用，以及默认清尺度和单独确认复制已知测试尺度。这里20×40m为受控素材的已知构造条件，不是桥梁实测尺寸；它验证功能契约，不证明隧道机位或真实尺度能够自动迁移。','plans/002-verification.md','docs/camera-templates.md')
h('2.6.3 人工标注 版本绑定与一对一评测',3)
p('人工核验已关闭段的完整录像，记录物理车辆过线、事件与已核验类别。标注绑定运行、来源、配置指纹、尺寸和revision，修改撤销完整声明、使旧结果过期。ID查询仅读取近邻实际CSV；完全漏检事件可无ID录入并计作漏报。','backend/annotations.py','docs/visual-evaluation.md')
p('标注草稿与完整真值有不同状态。只有操作者记录核验者、区间、已完整查看的类别并确认完整性，保存版本才可正式评测；未勾选类别表示尚未知，空数组只有在该类别完整核验后才表示没有目标。当前可视化界面不核验速度，speeds=[]与speed_state=unreviewed不能解释为确认没有速度对象。','docs/visual-evaluation.md','backend/annotations.py')
p('标注采用独立revision和运行配置指纹。修改姓名、区间、类别或条目撤销完整声明，保存时校验expected_revision，避免另一个编辑窗口覆盖新版本；评测与下载绑定精确保存版本。恢复本机草稿需显式选择，恢复后也撤销完整声明。本机缓存帮助防丢失，不等于服务真值或跨设备备份。','backend/annotations.py','docs/visual-evaluation.md')
p('暂停时的ID辅助查询仅从已有CSV读取0.5s以内最近的实际观测样本，不重新运行检测。它帮助建立物理车辆与当前段ID的映射，却不能确认系统检测正确；完全没有轨迹的事件应无ID记录，并计入漏报，不能用空ID匹配任意预测。独立播放器上一/下一帧按保存FPS近似步进，浏览器与可变帧率解码不保证精确逐帧定位，评测记录以显示并保存的原视频秒数为准。','docs/visual-evaluation.md','backend/annotations.py')
fig(7,'fig7-evaluation.png','人工评测结果区  受控流程验收，非真实精度实验',4.8)
p('事件按类型、映射ID和时间容差进行Hungarian一对一匹配，输出TP、FP、FN、precision和recall；默认容差0.15s可调整。无预测时precision、无真值时recall不可定义，未核验类别与速度保持unavailable。图7受控标签特意包含误报和未映射漏报，TP2/FP1/FN1只验证匹配契约，不代表真实交通报警性能。','backend/evaluation.py','docs/visual-evaluation.md','plans/002-browser-result.json','plans/002-verification.md')
p('正式评测先核对来源、配置、尺寸、区间和所评CSV的有效性，再进行计数及事件比较。区间计数由末端累计值减起点之前累计值，与人工过线条目数比较绝对误差；事件在类型、映射ID和容差共同约束下作一对一匹配，不能让一个机器报警重复匹配多个真值。输入损坏或所需CSV缺失时返回明确不可用，不把缺失文件视作零目标。','backend/evaluation.py','docs/visual-evaluation.md')
p('precision=TP/(TP+FP)，recall=TP/(TP+FN)。分母为零时指标不可定义，保存null或不可用；只核验计数时不计算事件指标。现有CLI还可对独立速度真值按ID和近邻实际观测时刻匹配，输出MAE、RMSE及匹配覆盖率，但当前没有可据此报告的真实测速结果。界面可完成记录和计算，不代表真值的人工完整性已经由软件自动证明。','backend/evaluation.py','docs/accuracy-evaluation.md')

newpage()
h('3 作品特色与创新',1)
p('以下特色属于当前系统设计与工程集成，不主张检测网络或基础跟踪算法创新。','backend/vision/detector.py','plans/002-verification.md')
h('3.1 观测与测量前提显式约束')
p('observed与标定区域门控避免将预测当测量，缺尺度则留空。回归测试及桥梁CSV的空世界坐标、空速度与零相关报警支持实现；明确车辆帧分母和窗口覆盖，使输出可解释，不构成精度证明。','backend/processor.py','backend/reports.py','tests/test_analysis_features.py','backend/data/runs/20261002-170236-c53990/vehicles.csv')
p('这一设计针对“界面存在数值便被理解为测量成立”的问题。实际实现把输入尺度、区域、观测状态和最少样本作为前提，未满足前提时保留空状态，并让原始导出与界面一致。可验证特色是数据含义明确和错误前提受到约束；不将采用已有最小二乘公式或预训练模型本身列作基础算法创新。','backend/processor.py','backend/traffic/speed_estimator.py','backend/exporter.py')
h('3.2 同帧展示与来源可追溯的复核')
p('图像与状态成对提交；截图、核验回放与人工记录分离，复核不调用实时定位。乱序解码测试及Range回放验收支持从二维状态返回原画面解释事件，截图不作为精度证据。','frontend/src/utils/snapshotState.ts','frontend/tests/snapshotState.test.mjs','backend/evidence.py','plans/2026-10-02-feature-verification.md')
p('同帧提交解决异步显示的工程一致性，独立回放解决历史复查与当前会话争用，不可变事实与独立复核记录解决人工判断覆盖机器历史的问题。三者形成从车辆状态到原始画面和条件记录的解释路径。已有浏览器验收观察到历史播放器定位时，后台当前分析仍继续完成；该证据支持操作隔离，不能推出端到端同步误差为零。','plans/002-verification.md','backend/evidence.py','frontend/src/utils/snapshotState.ts')
h('3.3 机位复用与人工评测版本绑定')
p('模板默认清尺度，复制实测尺度须单独确认；标注绑定指纹与版本，修改使旧结果失效。比例边界、修订冲突、部分类别和无ID漏报测试及浏览器保存重开验收支持实现，迁移效果仍待验证。','backend/camera_templates.py','backend/annotations.py','tests/test_camera_templates.py','tests/test_annotations.py','plans/002-verification.md')
p('机位模板减少重复几何输入，人工文档版本绑定防止旧结论冒充新标注。两者共同保留条件变更的可见性：几何复用仍需确认实测前提，标注变化必须重新完整核验与计算。贡献在于将复用效率与测量边界结合，并提供可追溯记录；它不是自动标定、自动生成真值或跨隧道迁移学习。','backend/camera_templates.py','backend/annotations.py','docs/visual-evaluation.md')
h('4 系统实现与实验验证',1)
h('4.1 实现环境 模型及演示素材')
p('本次环境为macOS 26.6 arm64、Python 3.13.0、FastAPI 0.136.3、PyTorch 2.11.0；前端为React 19、TypeScript、Vite、ECharts。已验证CPU，未实测NVIDIA GPU。主要素材为2160×3840、30FPS竖向桥梁录像，另有3840×2160横向素材验证方向与布局。','README.md','frontend/src/App.tsx','docs/review-smoke-2026-10-02.json','docs/twin-orientation.md')
p('后端以Python、FastAPI、OpenCV、PyTorch、NumPy和SciPy组合完成解码、检测、跟踪与几何统计；前端以React、TypeScript、Vite和ECharts提供工作台与趋势展示。本次核验还记录OpenCV 4.13.0.92、Torchvision 0.26.0及Ultralytics 8.4.32。这些版本说明本机验证条件，不构成跨平台安装或不同硬件性能承诺。','README.md','models.json')
p('竖向demo.mp4为主要机器运行证据，保存内容身份3a247cda14c4bfd01d7d6a3a，与本次重新计算的SHA256前24位一致。原视频共有256帧，元数据帧率30FPS；最后解码时间为8.5s，视频时长约8.533s，二者一个表示末帧PTS，一个含最后帧持续时间，不应认为记录丢失了0.033s。该运行未加入实测尺度。','backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','docs/review-smoke-2026-10-02.json')
h('4.2 功能流程与实际界面')
p('使用顺序为录像导入、道路配置与标定、分析及同帧展示、报告导出、独立回放和人工评测。图4、图6为真实桥梁界面，图7为受控验收。','plans/002-verification.md','plans/2026-10-02-feature-verification.md')
p('导入后先检查视频比例与已保存的道路轮廓，再确定有限计数线的位置和各车道合法方向。若没有真实道路尺寸，则直接在未标定状态演示检测、轨迹、过线和同帧孪生；不要为了出现速度而填入猜测尺度。具备可靠尺度后再设置四点区域，检查样本是否处于区域内，以及速度何时达到最少观测点条件。','README.md','docs/camera-templates.md','backend/processor.py')
p('分析结束后以统计段查看累计值和原始CSV，人工复核报警时检查触发时刻、目标ID、规则条件和邻域原视频。建立真值时应离开机器事件列表，独立核验整个区间并记录完全漏报。完成保存版本的评测后再下载真值与结果，保持run身份、配置和revision一致，避免将不同运行的ID用于比较。','docs/features-analysis.md','docs/visual-evaluation.md')
h('4.3 软件测试和运行验证')
p('2026年10月3日在隔离数据目录复测后端88项、前端27项均通过，并通过HTTP、可信来源空闲WebSocket和正常退出检查。表4区分本次复测、历史验收和实际输出。测试验证软件契约，不将其通过率换算成感知准确率。','tests/conftest.py','scripts/check_runtime.py','plans/002-verification.md')
p('后端回归覆盖道路几何有效性、有限计数线、测速与规则、导出报告、来源身份、模板、标注版本和评测边界；前端回归覆盖同帧状态提交、几何映射、人工输入及异步过期状态。本次测试使用隔离DATA并关闭缓存写入，未把测试尺度或构造真值写入用户历史运行目录。测试通过表示这些断言在当前环境成立，不等于真实交通场景覆盖完整。','tests/conftest.py','tests/test_analysis_features.py','tests/test_camera_templates.py','tests/test_annotations.py','frontend/tests/snapshotState.test.mjs')
p('运行检查启动隔离实例，核对HTTP响应、可信来源WebSocket空闲连接和正常关闭，结束后移除该实例。它验证服务基本可用与进程退出，不测实际摄像头断流重连，也不提供长时间负载或GPU性能结果。已有历史生产构建和浏览器证据与本次软件回归分别记录，本文没有重新声称完成全部历史界面流程。','scripts/check_runtime.py','plans/002-verification.md')
p('历史记录还覆盖独立接口22项、边界6项、320px和390px人工弹窗、草稿恢复及错误导入。部分验证使用200×400、10FPS、120帧受控素材与颜色检测器，事件标注包含刻意构造的FP和无ID FN；这类条件便于检查确定性行为，却不具有真实车辆检测和隧道异常识别的科学真值。','plans/002-verification.md')
table(4,'功能验证与实际运行证据',['验证层次','结果或材料','支持的结论及边界'],[
['本次软件回归','后端88通过；前端27通过','几何、规则、同步、版本与异常输入契约'],
['本次运行检查','HTTP / WebSocket / 正常退出通过','隔离合成素材；不证明交通精度'],
['历史实际浏览器','模板、报告、独立回放、人工评测','plans/002-verification.md；受控标签'],
['真实桥梁CSV','256帧；19次机器过线；28个ID','保存来源和配置；未经逐车真值核验'],
['观测样本口径','1074观测帧；1088预测帧','2162车辆帧行；有效观测速样本0'],
['既有CPU smoke','42.5FPS；相邻帧ID保留98.9%','单次离线吞吐和连续性；非精度']],[3.0,6.2,7.8],['plans/002-verification.md','scripts/check_runtime.py','docs/review-smoke-2026-10-02.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/vehicles.csv'])

newpage()
h('4.4 有真实真值支撑的评测状态')
p('尚无已保存的真实隧道或实测测速真值评测。运行段20261002-170236-c53990来源SHA256前24位与身份一致，配置完整、状态completed。图8绘制全部256行CSV：0.0–8.5s，累计过线19、最多在场13、确认轨迹ID共28个。曲线为机器输出，ID不等于物理车辆总数。','backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/vehicles.csv','plans/002-verification.md')
p('选择该段作为说明书的数据示例，是因为它保留原视频内容身份、分析配置、道路场景、标定状态和完整CSV，能够核对数字来源。较早运行和单次smoke用于说明历史进展，不混合为新的对照实验。当前没有另一个同真值、同区间并受控改变唯一条件的精度实验，也没有消融实验或提升百分比。','backend/data/runs/20261002-170236-c53990/run.json','plans/002-verification.md')
fig(8,'fig8-traffic.png','真实桥梁运行CSV交通曲线  全部统计帧，未经人工真值核验',4.3)
p('该段1074条实际观测车辆帧中有效速度样本0条，世界坐标、速度与加速度为空，机器事件0条；这验证未标定门槛，不能推断场景不存在异常。末帧每分钟流率134.12辆/min由19÷8.5×60短时折算，不是完整一分钟计数。既有smoke的42.5FPS为256帧离线处理吞吐；30FPS是原视频帧率，浏览器显示帧率受推送、解码与渲染影响，当前未独立测量。98.9%相邻帧ID保留率仅表示运行连续性。','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/events.csv','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','docs/review-smoke-2026-10-02.json')
p('2162行车辆明细由1074条实际观测和1088条预测构成，同一ID可在不同帧重复出现。19次机器过线由去重后的轨迹穿越产生，28个ID为整个区间内曾输出的确认轨迹，13为某一时刻最多同时在场轨迹。三个指标的时间范围和分母不同，不存在“19/28即检测准确率”这样的换算关系。事件0条首先与未标定门槛一致，无法证明原视频完全没有异常。','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/events.csv')
p('图7受控流程输出计数绝对误差0、TP2/FP1/FN1，对应precision与recall均为66.7%。这里的标签为工作流验收而构造，且包含专门检查未映射漏报的边界条目；数字只说明一对一匹配与界面下载一致，不列为真实报警准确率。它与图8的真实桥梁CSV属于不同证据层次，也不能据此选择最佳报警阈值。','plans/002-browser-result.json','plans/002-verification.md')
table(5,'真实桥梁运行结果及统计口径',['指标','实际记录','分母 条件与解释'],[
['统计区间与帧数','0.0–8.5s；256个统计帧','末帧PTS；30FPS原视频；来源身份一致'],
['确认轨迹与在场量','28个ID；最多同时13个','ID非物理车辆数；在场包含预测'],
['机器过线累计','19次','有限计数线；同ID同段去重'],
['车辆帧样本','1074观测+1088预测=2162行','重复ID跨帧计行；非2162辆车'],
['有效观测速样本','0/1074；覆盖率0%','未实测标定；均速与误差不可用'],
['机器事件与流率','0条；末帧134.12辆/min','报警门槛关闭；19÷8.5×60短时折算']],[3.8,5.5,7.7],['backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/events.csv'])
h('4.5 当前局限与后续验证计划')
p('现有桥梁录像不能覆盖隧道低照度、车灯反光、遮挡和相机震动；四点平面模型未补偿镜头畸变、道路高差及移动相机。需补充合法来源的固定隧道录像、机位与实测尺寸，独立完整标注各区间过线及事件，保存核验者、容差和版本，再报告计数绝对误差及事件TP/FP/FN。测速需独立参考测量和ID时间对应，报告匹配覆盖率、MAE、RMSE，不能仅报告匹配成功的样本。检测mAP和MOT类指标另需逐帧框及物理ID标注、相应导出；还需长时运行、GPU、干净安装及RTSP断流恢复验收。以上为后续方案，未作为已完成实验。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','plans/2026-10-02-implementation-status.md')
p('后续应先建立来源合法、场景可说明的连续视频集合，分开记录隧道入口与内部、明暗变化、遮挡和车流条件；保存机位、录像时段、文件身份及相机是否移动。建议由两名标注者独立核验并复查分歧，完整记录区间端点、物理车辆、系统ID映射、事件定义和容差。不以挑选少数效果良好片段代替整体报告。','docs/accuracy-evaluation.md')
p('计数实验应在同一来源、计数线和完整区间内比较机器计数与人工过线次数，报告分场景绝对误差，并解释ID分裂、有限线段漏穿越和端点采样误差。异常实验应明确首次满足规则的时刻和冷却口径，分别报告三类事件TP、FP、FN及分母；没有真实目标的类别报告样本缺失或不可定义，不制造高召回率。','docs/accuracy-evaluation.md','docs/visual-evaluation.md')
p('测速实验须取得独立参考速度，建立参考设备或人工距离时间记录与视频PTS、物理车辆、系统ID的对应。先报告有多少真值样本成功匹配，再给MAE和RMSE：只对成功测得样本平均会掩盖无输出和漏检。道路四点和参考测速不能都来源于系统自身，否则形成自我验证。真实测试规模应由取得的材料决定，本文不预填样本量或目标误差。','docs/accuracy-evaluation.md')
p('软件层面还需验证长时写入、视频来源缺失、磁盘空间不足、RTSP断流恢复、不同硬件推理和干净安装。新增实时流入口必须同时验收有效时间、失败状态与恢复行为；当前底层VideoCapture支持流字符串不等于产品已具备完整接入能力。多摄像头扩展也需重构单分析会话的资源管理与隔离，不能简单增加一个下拉选项即称已完成。','backend/vision/video_reader.py','backend/service.py','plans/2026-10-02-implementation-status.md')
table(6,'已完成工作与待验证事项',['领域','已完成及证据','待补材料或验收'],[
['检测与跟踪','预训练适配、Kalman/IoU；测试与真实smoke','逐帧框与物理ID真值；mAP/MOT评测'],
['道路测量','四点计算、时间拟合、未标定留空','机位与尺寸实测；独立参考速度'],
['统计与事件','有限线段、窗口、规则、证据保存','真实完整过线及三类事件标签'],
['二维展示与复核','同帧提交、回放、模板、版本评测','代表性隧道场景与机位复用验证'],
['运行与推广','本机CPU、单会话、HTTP/WS检查','RTSP、长时、GPU、安装与容量验证']],[3.4,6.7,6.9],['docs/accuracy-evaluation.md','docs/visual-evaluation.md','plans/002-verification.md','plans/2026-10-02-implementation-status.md'])
h('5 应用前景与推广条件',1)
p('系统可用于监测教学、离线录像分析和事件人工复查。推广需取得固定机位与实测尺度，完成代表性隧道验证、现场阈值配置、来源授权、存储和故障处置。当前无落地案例、多摄像头并发或多用户权限，生产应用还需流接入、容量和持续运行保障。模型与视频许可须核实，不能承诺商业化直接使用。','models.json','backend/service.py','docs/camera-templates.md','plans/002-verification.md')
p('近期可应用于教学演示与离线分析：学生或评委从原视频观察车辆，理解检测、跟踪、标定、时间拟合和统计之间的关系，并通过导出和独立回放复核结论。当前成果适合解释算法与系统协同的实现，应用价值首先来自可理解、可检查的记录链，而不是未经验证的检测精度或自动执法能力。','README.md','docs/features-analysis.md')
p('推广到固定隧道机位时，应先完成场景验证与尺度确认，再设置车道、方向、统计断面和报警阈值。机位模板可减少同机位重复配置，条件对照可辅助观察模型或参数输出变化，人工评测可检查真实错误。三者只有在来源与真值可靠时，才能支持后续选型；目前不能承诺阈值无需调整或不同隧道直接复用。','docs/camera-templates.md','docs/accuracy-evaluation.md')
p('生产使用还涉及授权、模型许可、隐私与数据保留、事件处置责任和服务可靠性。本文只提出这些推广条件，不编写尚未发生的合作、部署、经济收益或落地案例。最终竞赛提交前还应由团队核实人员信息、素材和权重使用条件、赛事实际模板及真实性声明；本文不沿用参考作品人员，也不代签承诺书。','models.json','plans/2026-10-02-implementation-status.md')
h('参考文献',1)
refs=[
('[1] Veres M, Moussa M. Deep Learning for Intelligent Transportation Systems: A Survey of Emerging Trends[J]. IEEE Transactions on Intelligent Transportation Systems, 2020, 21(8): 3152-3168. DOI: 10.1109/TITS.2019.2929020.','https://doi.org/10.1109/TITS.2019.2929020'),
('[2] Sochor J, Juránek R, Herout A. Traffic Surveillance Camera Calibration by 3D Model Bounding Box Alignment for Accurate Vehicle Speed Measurement[J]. Computer Vision and Image Understanding, 2017, 161: 87-98. DOI: 10.1016/j.cviu.2017.05.015.','https://doi.org/10.1016/j.cviu.2017.05.015'),
('[3] Sochor J, Juránek R, Špaňhel J, et al. Comprehensive Data Set for Automatic Single Camera Visual Speed Measurement[J]. IEEE Transactions on Intelligent Transportation Systems, 2019, 20(5): 1633-1643. DOI: 10.1109/TITS.2018.2825609.','https://doi.org/10.1109/TITS.2018.2825609'),
('[4] dronefreak. visdrone-yolo11n 模型卡[EB/OL]. [2026-10-03].','https://huggingface.co/dronefreak/visdrone-yolo11n'),
('[5] Ultralytics. Ultralytics YOLO11 官方文档[EB/OL]. [2026-10-03].','https://docs.ultralytics.com/models/yolo11/'),
('[6] Torchvision. ssdlite320_mobilenet_v3_large 官方文档[EB/OL]. [2026-10-03].','https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.detection.ssdlite320_mobilenet_v3_large.html'),
('[7] Zhu P, Wen L, Bian X, et al. Vision Meets Drones: A Challenge[EB/OL]. 2018. arXiv:1804.07437.','https://arxiv.org/abs/1804.07437'),
('[8] Bewley A, Ge Z, Ott L, et al. Simple Online and Realtime Tracking[C]. ICIP, 2016. DOI: 10.1109/ICIP.2016.7533003.','https://doi.org/10.1109/ICIP.2016.7533003'),
('[9] OpenCV. Geometric Image Transformations[EB/OL]. [2026-10-03].','https://docs.opencv.org/4.13.0/da/d54/group__imgproc__transform.html'),
]
from docx.opc.constants import RELATIONSHIP_TYPE as RT
for text,url in refs:
    pa=doc.add_paragraph(text);pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(1)
    for r in pa.runs:r.font.size=Pt(10.5)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz);rr.append(rp)
    tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=' 在线来源';rr.append(tt);link.append(rr);pa._p.append(link)


path=OUT/'基于视频感知的隧道交通数字孪生系统作品说明书.docx';doc.save(path)
(OUT/'evidence/manuscript.md').write_text('\n\n'.join(draft),encoding='utf-8')
with (OUT/'evidence/claims.csv').open('w',newline='',encoding='utf-8-sig') as f:
    wr=csv.DictWriter(f,fieldnames=list(claims[0]));wr.writeheader();wr.writerows(claims)
(OUT/'evidence/displays.json').write_text(json.dumps(displays,ensure_ascii=False,indent=2))
print(path)
