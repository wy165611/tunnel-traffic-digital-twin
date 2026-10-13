from pathlib import Path
import csv,json,hashlib,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.5)
sec.left_margin=sec.right_margin=Cm(2)
sec.footer_distance=Cm(1.25)
for name,size in [('Normal',12),('Title',18),('Heading 1',15),('Heading 2',12.5),('Heading 3',12),('Caption',10.5)]:
    s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'STSong')
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        s._element.rPr.rFonts.attrib.pop(qn('w:'+key),None)
    if s._element.pPr is not None:
        for border in list(s._element.pPr.findall(qn('w:pBdr'))):s._element.pPr.remove(border)
    pf=s.paragraph_format; pf.line_spacing=1.18; pf.space_before=Pt(0); pf.space_after=Pt(4)
    if name.startswith('Heading'): pf.space_before=Pt(8); pf.space_after=Pt(4); s.font.bold=True
    if name=='Normal': pf.first_line_indent=Pt(24)
    else: pf.first_line_indent=Pt(0)
    if name=='Title':s.font.bold=True
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.core_properties.title='基于视频感知的隧道交通数字孪生系统作品说明书'
doc.core_properties.author=''; doc.core_properties.subject='智能交通竞赛作品说明书'
doc.core_properties.keywords='隧道交通 视频感知 二维数字孪生 本科科研 实验验证'
doc.core_properties.comments=''
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
r=footer.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)

manifest=json.loads((OUT/'evidence/source_manifest.json').read_text())
REFIDS={1:'E088',2:'E076',3:'E077',4:'E078',5:'E079',6:'E080',7:'E081',8:'E082',9:'E083',10:'E084'}
E={x['path']:x['evidence_id'] for x in manifest['sources']}
claims=[]; draft=[]; displays=[]
def evidence(*paths): return [E[p] for p in paths]
def record(text,paths):
    ids=evidence(*paths) if paths else []
    for match in re.finditer(r'\[(\d+)(?:-(\d+))?\]',text):
        lo=int(match[1]);hi=int(match[2] or match[1]);ids.extend(REFIDS[n] for n in range(lo,hi+1) if n in REFIDS)
    ids=sorted(set(ids))
    c=f'C{len(claims)+1:03d}'
    claims.append({'claim_id':c,'claim_text_sha256':hashlib.sha256(' '.join(text.split()).encode()).hexdigest(),'evidence_ids':','.join(ids),'verification_status':'human_pending','assistant_checked_on':'2026-10-06'})
    draft.append(f'{text} [claim:{c}] [evidence:{",".join(ids)}]')
def add_cited_text(pa,text):
    for part in re.split(r'(\[\d+(?:[-,]\d+)*\])',text):
        run=pa.add_run(('\u2060'+part) if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',part) else part)
        if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',part):
            run.font.superscript=True; run.font.size=Pt(9)
def p(text,*paths,small=False):
    pa=doc.add_paragraph()
    if pa._p.getprevious() is not None and pa._p.getprevious().tag==qn('w:tbl'):pa.paragraph_format.space_before=Pt(6)
    add_cited_text(pa,text)
    if small:
        pa.paragraph_format.first_line_indent=Pt(0)
        for r in pa.runs:r.font.size=Pt(10.5)
    record(text,paths)
    if text.startswith(('本作品以异常停车','本作品的特色来自交通需求')):pa.paragraph_format.keep_together=True
    if text.startswith('H为3×3'):pa.paragraph_format.keep_with_next=True
    return pa
def h(text,level=2):
    doc.add_paragraph(text,'Heading '+str(level));draft.append('## '+text)
def newpage():pass
def fig(n,name,caption,height):
    im=OUT/'assets'/name
    from PIL import Image
    w,hh=Image.open(im).size; width=min(17,height*w/hh)
    pa=doc.add_paragraph(); pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
    pa.alignment=WD_ALIGN_PARAGRAPH.CENTER; pa.paragraph_format.keep_with_next=True
    pic=pa.add_run().add_picture(str(im),width=Cm(width)); pic._inline.docPr.set('descr',caption)
    cp=doc.add_paragraph(f'图{n}  {caption}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    displays.append({'number':n,'kind':'figure','asset':name,'caption':caption,'processing':'local diagram or screenshot crop; no fact alteration'})
def table(n,title,headers,rows,widths,paths):
    cp=doc.add_paragraph(f'表{n}  {title}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for col,w in zip(tb.columns,widths):col.width=Cm(w)
    for cell,w in zip(tb.rows[0].cells,widths):cell.width=Cm(w)
    for cell,text in zip(tb.rows[0].cells,headers):cell.text=text
    for row in rows:
        cells=tb.add_row().cells
        for cell,w,text in zip(cells,widths,row):cell.width=Cm(w);cell.text=str(text)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'183E59' if ri==0 else ('F1F6F8' if ri%2==0 else 'FFFFFF'));tcpr.append(sh)
            for pa in cell.paragraphs:
                pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0);pa.paragraph_format.line_spacing=1.08
                if ci==0 or ri==0 or (n==5 and ci==1):pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if ri==0:pa.paragraph_format.keep_with_next=True
                raw=pa.text
                pa.clear();add_cited_text(pa,raw)
                for r in pa.runs:
                    if not r.font.superscript:r.font.size=Pt(11)
                    r.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    record('；'.join([title]+[' / '.join(str(x) for x in row) for row in [headers]+rows]),paths)
    displays.append({'number':n,'kind':'table','title':title,'native_word_table':True,'evidence':evidence(*paths)})
def eq(prefix,numerator,denominator,suffix):
    pa=doc.add_paragraph();pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0)
    m=OxmlElement('m:oMath')
    def mr(t):
        r=OxmlElement('m:r'); x=OxmlElement('m:t');x.text=t;r.append(x);return r
    m.append(mr(prefix));f=OxmlElement('m:f')
    for tag,t in [('num',numerator),('den',denominator)]:
        z=OxmlElement('m:'+tag);z.append(mr(t));f.append(z)
    m.append(f);m.append(mr(suffix));pa._p.append(m)

pa=doc.add_paragraph('基于视频感知的隧道交通数字孪生系统','Title');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['参赛团队：待补充    指导教师：待补充','学校及院系：待补充    所在城市及邮编：待补充']:
    pa=doc.add_paragraph(text);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
pa=doc.add_paragraph('作品内容简介','Heading 1');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('本作品面向隧道交通监测中视频画面难以转化为直观交通状态、异常提示缺少复核依据和分析结果难以追溯的问题，设计并实现视频感知与二维道路数字孪生系统。作品将车辆识别、道路配置、交通统计、同步展示和人工复核组织为完整工作流程，使使用者能够从原视频了解车辆位置、车道和行驶变化，并回到原始记录检查系统判断。项目采用现成预训练模型，研究重点在于交通需求分析、系统方案设计和可验证的工程实现。其特色包括原视频与孪生同帧对应、测量前提明确约束、报警证据独立复核，以及机位模板和版本化人工评测。已有桥梁俯拍录像运行、软件回归和浏览器验收支撑当前功能，真实隧道适应性与测速精度仍需独立实测。作品当前为本机单分析会话和二维展示，未自行训练或改进检测网络；缺少可靠道路实测尺度时，速度、加速度及相关报警保持不可用。现阶段成果可作为交通监测教学、离线录像研究和事件复查的原型，后续具备隧道素材、标定及真实评测条件后，可进一步研究固定机位的辅助监测应用。','docs/features-analysis.md','docs/visual-evaluation.md','backend/processor.py','backend/service.py','plans/002-verification.md')
p('关键词：隧道交通监测；视频感知；二维数字孪生；系统设计；人工复核',small=True)

h('1 设计背景与研究目标',1)
h('1.1 从观看交通视频到理解交通状态')
p('隧道交通监测的基本任务，是及时了解道路上的车辆和疑似异常，并为人工判断提供可定位、可复查的信息。世界道路协会PIARC指出，多幅监控画面和其他值守任务会增加人工观察的难度，自动事件提示可辅助发现停车等情况，运营人员仍需确认事件性质和位置[1]。这一需求构成了本作品选择视频感知与人工复核相结合路线的出发点。','https://tunnelsmanual.piarc.org/en/equipment-systems-monitoringincident-detection/automatic-incident-detection')
p('原视频保留了现场细节，但观察者需要自行识别车辆、判断道路位置、记住前后状态并统计过线。本作品将这些工作转化为具体研究问题：能否把视频中的车辆组织为连续状态；能否把车辆变化对应到道路平面；能否让统计和异常提示具有可检查的依据。这样的选题既联系交通应用，又能够形成范围清楚、可逐步验证的原型。','README.md','docs/features-analysis.md','docs/accuracy-evaluation.md')
p('本作品关注信息的组织与解释。车辆框帮助找到目标，轨迹帮助观察连续变化，车道和计数线帮助明确统计对象，二维孪生帮助理解车辆在道路上的相对位置。报告和独立回放进一步把一次观察变成可以保存、复查和比较的记录。这些功能之间存在依赖关系，因此作品需要形成一个完整系统，才能回答实际使用中的问题。','backend/processor.py','backend/reports.py','backend/evidence.py','frontend/src/components/DigitalTwin.tsx')
p('这里的数字孪生具体指视频状态在二维道路视图中的同步映射。它服务于交通状态展示和解释，当前不包含三维场景重建、车辆动力学仿真或交通控制推演。采用这样的范围，能够把有限的开发与验证工作集中在视频、位置和证据的一致性上，也使作品目标与已有成果保持对应。','frontend/src/components/DigitalTwin.tsx','docs/twin-orientation.md','backend/service.py')
h('1.2 需求分析与需要解决的问题')
p('从使用过程出发，作品将需求归纳为看清状态、形成统计和检查判断三个相互连接的方面。看清状态要求视频目标与道路位置对应；形成统计要求明确计数线、车道和时间范围；检查判断要求能够找到触发时刻、分析条件和原始画面。图1先概括从视频观察到形成可复查记录的需求关系，表1进一步对应到当前功能，说明每个模块为何需要建设。','docs/features-analysis.md','docs/visual-evaluation.md')
fig(1,'fig1-needs.png','作品要解决的观察与研究需求  根据项目需求绘制',3.1)
table(1,'使用需求与作品功能的对应关系',['使用需求','需要回答的问题','当前实现'],[
['交通状态观察','车辆在哪里，前后如何变化','视频与二维孪生同帧展示，轨迹与车道标识'],
['断面交通统计','某段时间有多少轨迹过线','有限计数线、类别统计及10/30/60秒车道窗口'],
['疑似异常提示','提示由什么变化触发','满足前提时的停车、急减速和逆行规则'],
['结果核查','能否返回原画面检查','报警截图、独立原视频回放和人工复核'],
['研究记录保存','能否复查条件和重新评价','历史报告、CSV/HTML、机位模板与版本化评测']],[3.1,5.2,8.7],['docs/features-analysis.md','docs/visual-evaluation.md','docs/camera-templates.md'])
p('这些需求还包含一个容易忽略的问题：系统有输出，并不意味着输出已经可信。例如检测框短暂消失时，预测位置可使画面连续，但不能当作新的实际测量；没有道路真实尺寸时，可以观察轨迹，却不能据此给出真实车速。作品把这类前提纳入设计，使“有数据”和“有可靠测量”在使用中能够被区分。','backend/processor.py','backend/traffic/speed_estimator.py','docs/accuracy-evaluation.md')
p('异常提示同样需要可解释的工作流程。一个停车提示只能说明在设定条件下持续低速，不能直接判断事故原因；人工确认某条报警，也不能证明系统发现了全部异常。因此，作品同时提供报警复核和全区间人工标注，前者用于解释单条结果，后者用于发现漏报并评价整体表现。','backend/events/rules.py','backend/evidence.py','backend/annotations.py','docs/visual-evaluation.md')
p('需求分析也决定了功能之间的优先关系。首先建立视频与车辆状态的对应，随后定义车道和统计断面，再把变化转化为可复查提示。若只有模型输出而没有道路对象，统计难以解释；若只有报警而没有来源和时间记录，判断难以复核。因而，系统组织本身成为本作品需要设计和验证的研究对象。','backend/processor.py','docs/features-analysis.md')

h('1.3 相关研究与本作品的切入点')
p('智能交通领域已形成较丰富的视频分析方法。Veres和Moussa的综述讨论了深度学习在交通任务中的应用及场景、时空特征带来的挑战[2]。对本作品而言，已有检测模型提供了可使用的基础组件，项目仍需解决模型如何与道路配置、状态展示、统计和人工判断配合的问题。','https://ieeexplore.ieee.org/document/8771378','backend/processor.py')
p('单目视频测速研究也说明了尺度和独立验证的重要性。Sochor等研究了利用三维车辆模型进行摄像机标定的方法[3]，并建立了含独立速度参考测量的数据集[4]。本作品采用人工四点与实测道路尺寸的路线，没有实现该文献中的三维自动标定，也未使用其数据集作为本项目实验结果。相关研究用于确定测量前提和验证思路。','https://arxiv.org/abs/1702.06451','https://arxiv.org/abs/1702.06441','backend/calibration/homography.py','docs/accuracy-evaluation.md')
p('作品的研究切入点是把成熟组件组织为可使用、可检查的交通分析原型。在现有预训练模型基础上，设计同帧展示、测量约束、证据复核和人工评测等功能，研究它们怎样共同支撑结果解释。创新评价应针对这些具体设计及其验证依据，不把使用现成网络或已有跟踪方法本身作为原创算法成果。','models.json','backend/vision/detector.py','docs/features-analysis.md','docs/visual-evaluation.md')
h('1.4 研究目标与阶段范围')
p('作品的直接目标，是完成从录像输入到二维展示、统计归档和人工核查的工作流程，并证明各项功能在明确条件下能够运行。进一步的研究目标，是在取得真实隧道素材、可靠道路尺度与独立人工真值后，评价计数、测速和事件判断的表现。两个目标属于不同阶段：当前实现与功能验收已经有材料支撑，真实交通精度仍待实测。','plans/002-verification.md','docs/accuracy-evaluation.md')
p('目前主要真实素材为桥梁俯拍录像。选择它进行原型验证，是现有材料条件下检验车辆感知、道路映射和界面工作流程的可行路径；桥梁运行结果不能直接说明系统适应隧道。隧道目标保留在选题与后续研究中，已完成场景则在实验部分按桥梁素材说明。','README.md','plans/002-real-video-verification.json','docs/twin-orientation.md')
p('当前系统在本机运行，每次以一个分析会话处理录像。底层读取接口保留RTSP流字符串的支持，但现有入口和验收以文件分析为主，尚未完成实时摄像头接入验收。多摄像头并发、跨镜头身份关联、三维仿真及生产部署也未作为本阶段成果。限定范围有助于对作品进行具体评价，并为下一阶段研究保留清楚的问题。','backend/vision/video_reader.py','backend/service.py','plans/2026-10-02-implementation-status.md')

h('2 作品方案与系统设计',1)
h('2.1 整体方案及使用过程')
p('系统由视频感知、道路与交通分析、二维工作台以及历史与证据管理四部分组成。后端分析并保存记录，前端负责配置、展示和人工操作。图2依据源码绘制，表示当前本机单分析会话的联系。','backend/service.py','backend/processor.py','backend/main.py','frontend/src/App.tsx')
fig(2,'fig2-architecture.png','本机系统架构与数据联系  依据当前源码绘制',4.9)
p('对使用者而言，主要过程是导入录像、配置道路、观察分析、查看报告和核查结果。道路配置包括有效区域、车道、计数线及合法方向；如具备实测尺寸，再进行四点标定。分析时可同时查看原视频与二维孪生；分析结束后，以保存的统计段读取报告、导出数据，并通过独立播放器和人工评测检查结论。这一过程把即时观察与后续核查联系起来。','README.md','docs/features-analysis.md','docs/camera-templates.md','docs/visual-evaluation.md')
p('这种组织方式使界面操作与研究记录相连。更换视频、定位或调整关键条件后，系统形成新的统计段，保存对应来源和配置，避免把不同条件的结果混为一组。当前播放会话负责即时处理，历史统计段负责结果追溯，二者承担不同职责。','backend/service.py','backend/exporter.py','backend/reports.py')
h('2.2 基础组件的选择与自主建设内容')
p('车辆检测采用现成预训练模型，支持VisDrone YOLO11n、COCO YOLO11n和SSDLite。默认模型来自dronefreak发布的VisDrone微调权重[5]；另外两类来自Ultralytics与Torchvision官方发布[6-7]。VisDrone研究面向无人机视角的检测与跟踪任务[8]，与当前俯拍演示存在视角关联，但是否适合具体隧道仍需实景检验。','models.json','backend/vision/detector.py')
table(2,'预训练模型来源与方案选择',['可选模型','来源','本项目选择依据'],[
['VisDrone YOLO11n','dronefreak发布的微调权重[5]','当前俯拍演示默认；已有素材抽查支持选用'],
['COCO YOLO11n','Ultralytics官方预训练权重[6]','通用视角备选；新场景需要重新评价'],
['SSDLite MobileNetV3','Torchvision官方COCO权重[7]','轻量备选；当前俯拍抽查存在漏检']],[4.2,5.7,7.1],['models.json','backend/vision/detector.py','README.md'])
p('项目早期在现有俯拍素材上发现SSDLite漏检较多，随后把VisDrone YOLO11n作为默认选择。这体现了依据实际输入检查组件适用性的方案调整，但已有抽查没有统一人工真值，不能写成严格的模型精度对比实验。后续模型评价应使用相同视频区间和同一套真值，同时考察错误情况与本机处理成本。','README.md','docs/accuracy-evaluation.md','docs/review-smoke-2026-10-02.json')
p('项目建设内容包括把不同检测器接入统一处理流程，配置道路与统计对象，连接视频和二维状态，设计历史报告、报警证据、独立回放、模板以及人工标注与评测。这些工作使检测模型成为交通研究原型的一部分。这些建设成果通过源码、运行记录和功能验收留下可检查的依据。','backend/processor.py','backend/reports.py','backend/evidence.py','backend/camera_templates.py','backend/annotations.py','frontend/src/App.tsx')
p('采用成熟模型可以把本阶段工作集中到需求分析和系统验证。与此相应，作品必须说明权重来源、适用场景和使用边界，不能把发布者训练成果写成团队训练成果。模型文件在加载前进行摘要校验，这有助于核对文件一致性，但不证明其检测精度。','models.json','backend/vision/weights.py','backend/vision/detector.py')
h('2.3 从识别车辆到连续观察')
p('单帧检测只能说明某一时刻出现了哪些目标。为了观察前后变化，系统使用SORT风格的跟踪过程：先预测已有轨迹的大致位置，再按检测框重叠程度进行一对一匹配，维持轨迹ID。所采用的Kalman预测与Hungarian匹配属于已有方法[9]，作品重点是将连续轨迹接入道路统计和复核流程。','backend/vision/tracker.py','backend/processor.py')
p('系统区分实际检测关联的位置和缺少本帧检测时的预测位置。预测有助于展示短暂遮挡或检测间隔中的连续状态；测速、过线和异常规则只由实际观测更新。保存数据也保留这一区别，使报告可以使用正确的样本范围。这样的约束有助于避免把平滑的展示效果误认为连续发生的可靠测量。','backend/processor.py','backend/reports.py','backend/exporter.py')
p('轨迹ID是一次分析中的关联结果，不是车辆永久身份。遮挡、框偏移或重新进入画面都可能导致ID变化；同一物理车辆也可能形成多个轨迹。因此，确认轨迹数、同时在场轨迹数与人工车辆总数不能直接等同。项目保留ID用于观察和复核，其身份正确性仍需要独立标注检验。','backend/vision/tracker.py','backend/traffic/flow_counter.py','docs/accuracy-evaluation.md')
h('2.4 道路位置与统计对象的确定')
p('道路标定把图像位置对应到道路平面。系统取检测轨迹框的底边中心近似车辆接触道路的位置，由四个道路角点与对应实测长宽进行平面映射[10]。像素坐标表示画面位置，米制坐标还需要可靠实测尺度。道路映射本身不能证明尺寸已经实测。','backend/calibration/homography.py','backend/processor.py')
p('这种人工配置适合在固定机位下明确道路区域，方法也容易解释和复查。其代价是依赖角点选择、真实尺寸和道路近似平面的成立。程序能检查四边形与计算是否有效，却不能自动证明输入尺寸已经实测；也没有根据无人机高度、焦距或平均车长自动生成真实尺度。','backend/calibration/homography.py','docs/camera-templates.md','docs/accuracy-evaluation.md')
p('在可靠标定后，系统利用车辆位置随实际视频时间的变化，采用短窗口拟合并平滑计算速度，以减少单帧位置波动。时间采用解码得到的实际时间戳，不能简单用处理耗时替代。未可靠标定时，世界坐标、速度与加速度留空，相关报警关闭；这项约束贯穿界面、统计和导出。','backend/traffic/speed_estimator.py','backend/vision/video_reader.py','backend/processor.py','backend/exporter.py')
p('交通统计通过有限计数线定义观察断面，以实际观测轨迹穿越线段进行计数，同一轨迹ID在同一统计段内只计一次。车辆类别和过线位置对应的车道一起保存，便于查看不同类型与不同车道的机器输出。尚未划分的车道保留为未划分项，不用推测标签补齐。','backend/traffic/flow_counter.py','backend/processor.py','backend/exporter.py')
p('累计计数用于回答整个分析段的机器过线次数，10、30、60秒窗口用于查看最近一段时间的变化。窗口未覆盖完整时长时应说明实际覆盖秒数；每分钟流率可能是短时折算，而非实际连续观察一分钟。车道记录中同时展示的在场数量和均速属于采样时刻的状态，不能误称为整个窗口的平均值。','backend/processor.py','backend/reports.py','docs/features-analysis.md')
h('2.5 二维孪生的展示作用')
p('二维孪生把车辆状态放在道路平面上，目的是降低观察原画面时的透视干扰，帮助使用者理解位置、方向和车道关系。界面保留原视频，使抽象状态能够与实际画面相互参照。有可靠尺度时使用米制道路位置；未标定时按原图比例展示，并明确测量不可用。','frontend/src/components/DigitalTwin.tsx','frontend/src/utils/twinGeometry.ts','docs/twin-orientation.md')
p('系统通过FastAPI与WebSocket发布包含图像、车辆、轨迹、统计和事件的同帧快照，前端等待对应图像解码后一起提交，过期响应不会覆盖较新的状态。图3使用后续浅色界面验收的真实截图，分别放大原视频与二维道路视图，可核对同帧ID 22—29。该截图是功能展示证据，不是测速或检测精度证明。','backend/main.py','backend/service.py','frontend/src/utils/snapshotState.ts','frontend/tests/snapshotState.test.mjs')
fig(3,'fig3-twin.png','真实桥梁视频与二维孪生  同帧截图及ID 22放大，未标定',12.7)
p('车辆朝向依据轨迹位移形成显示方向，不是直接识别车头姿态。孪生视图可以旋转到便于阅读的道路方向，但不改变原视频和原始记录。图3使用竖向桥梁演示素材，来自2026年10月3日后续界面验收；它与后文CSV示例属于不同分析会话，不能仅凭相似指标认定是同一统计段。','frontend/src/components/DigitalTwin.tsx','frontend/src/utils/twinGeometry.ts','docs/twin-orientation.md','docs/linear-design/desktop-full.jpg','docs/linear-design/README.md')
h('2.6 异常提示与人工复核')
p('本作品以异常停车、急减速和逆行作为已实现的疑似事件类型，采用可以解释的轨迹与速度规则。默认停车条件为速度低于3km/h并持续至少3秒；急减速条件为速度变化率低于−4m/s²并持续至少0.5秒；逆行结合车道合法方向、速度大于3km/h和至少1秒的持续时间判断。这些规则只有在有效标定和实际观测等前提满足时启用。','backend/events/rules.py','config.yaml','backend/processor.py')
p('规则适合解释系统为何提示某辆车，但不理解事故原因或道路管理语义。排队可能被判为停车，检测抖动可能影响急减速，车道方向配置错误可能影响逆行。当前规则不是经过真实隧道样本优化的统一阈值；使用者应结合原视频和场景信息复核，并通过真实标注研究误报与漏报。','backend/events/rules.py','docs/accuracy-evaluation.md','docs/visual-evaluation.md')
fig(4,'fig4-evidence.png','疑似事件的证据保存与人工判断  依据当前实现绘制',3.25)
p('系统触发事件后保存截图、时刻、轨迹、规则与标定条件，人工判断另存为复核记录。独立播放器核对视频来源身份后回放原始片段，不改变当前分析位置，也不覆盖机器事件事实。图4展示从提示到解释的路径，使评委能够理解作品如何支持人工判断，而不仅是展示一个报警标记。','backend/evidence.py','docs/features-analysis.md','plans/2026-10-02-feature-verification.md')
h('2.7 历史分析与研究资料管理')
p('历史报告保存分析范围、累计机器过线、事件分类和有效观测速样本情况，并提供CSV、HTML与趋势图。它使一次演示可以留下可核对的数据，而不是在页面关闭后失去依据。报告界面示例见第4.2节图6A，用于核对统计区间和累计机器计数。','backend/reports.py','backend/exporter.py','docs/features-analysis.md')
p('条件对照先检查两个统计段的视频身份、时间范围、道路几何和尺度等前提，再列出模型或参数差异。这样的设计有助于发现比较对象是否一致，为后续实验提供记录基础。当前对照显示的是机器输出差异，只有另有同一套独立真值时，才能讨论哪种条件具有更好的准确性。','backend/reports.py','docs/features-analysis.md')
p('固定机位模板保存已确认的道路配置，减少重复输入。模板按视频比例检查并映射几何，默认清除实测尺度；复制尺度需要单独确认。相同比例不表示相同机位，因此模板是一项配置复用工具，不能代替机位核验和重新测量。','backend/camera_templates.py','docs/camera-templates.md','plans/002-verification.md')
p('人工评测支持完整区间的过线和事件标注，并把标注与视频来源、配置及版本绑定。修改标注后，旧核验声明和旧评测结果不能继续代表新版本。系统按一对一关系匹配事件，未映射到机器轨迹的真实事件仍可记录为漏报。这些功能为后续实验准备工具，真实标签的完整性和正确性仍需人工负责。','backend/annotations.py','backend/evaluation.py','docs/visual-evaluation.md')

h('3 作品创新与研究特色',1)
h('3.1 创新定位与评价依据')
p('本作品的特色来自交通需求驱动的系统设计，而不是自研检测网络。创新评价以具体问题、实现方式和已保存的验证材料为依据。表3概括四项设计特色，其价值在于让车辆状态、测量前提、原始证据和人工评价形成联系；是否优于其他系统，还需要同条件比较，本文不主张首创或性能领先。','models.json','docs/features-analysis.md','docs/visual-evaluation.md','plans/002-verification.md')
table(3,'作品设计特色及验证依据',['设计特色','针对的问题','依据与当前边界'],[
['视频与孪生同帧对应','抽象车辆状态难与原画面对照','真实双视图和乱序提交测试；未量化端到端延迟'],
['测量前提显式约束','预测位置或猜测尺度被当成测量','观测门控、未标定留空；不等于测量精度已验证'],
['事件证据独立复核','报警缺少解释，历史复查干扰当前分析','截图事实与复核分存、独立回放验收'],
['复用与评测版本管理','重复配置及旧标注结果误用','模板、修订冲突和重开保存验收；实景迁移待测']],[4.2,5.6,7.2],['frontend/tests/snapshotState.test.mjs','backend/processor.py','backend/evidence.py','backend/annotations.py','plans/002-verification.md'])
p('图5说明这些特色怎样共同支撑研究：同帧对应帮助解释状态，测量约束明确数据可用前提，证据与版本保存结论形成条件。三个方面共同服务于复查和重复研究，图中关系是对已实现设计的归纳，不代表已证明用户效率或真实交通精度提升。','frontend/src/utils/snapshotState.ts','backend/processor.py','backend/evidence.py','backend/annotations.py')
fig(5,'fig5-design.png','作品特色之间的联系  源码与功能归纳图',4.05)
h('3.2 从双视图展示形成可解释的交通观察')
p('只展示检测框时，使用者仍需在透视画面中理解车辆与车道的关系；只展示孪生时，又可能看不到模型遗漏或错误。因此作品保留同帧原视频和二维状态，使抽象表达具有可回查的现场依据。车辆、轨迹、车道和事件同时呈现，构成了交通状态解释的基本方式。','frontend/src/components/VideoView.tsx','frontend/src/components/DigitalTwin.tsx','frontend/src/utils/snapshotState.ts')
p('这一特色体现为信息之间的对应关系。后端按整帧形成结果，前端把图像解码和状态提交配对，避免较晚返回的旧图像或旧状态覆盖新结果。已有同步回归和真实截图支持它在当前流程中的实现。其应用意义是让使用者更容易核对同一时刻的状态；实际阅读效率是否提高，需要后续用户研究，不在本次结果中给出推测百分比。','backend/processor.py','frontend/src/utils/snapshotState.ts','frontend/tests/snapshotState.test.mjs','output/playwright/twin-horizontal.png')
p('在事件查看任务中，使用者可以先从道路视图定位目标，再以相同ID返回原视频观察前后变化，最后检查规则和统计。界面把这些信息放在同一分析上下文中，使解释无需脱离当前车辆与时刻。这里描述的是已实现功能可支持的操作方式；它为后续评价观察效率和理解难度提供了明确任务。','frontend/src/components/VideoView.tsx','frontend/src/components/DigitalTwin.tsx','backend/evidence.py','docs/features-analysis.md')

h('3.3 把可信使用条件纳入功能设计')
p('交通系统不仅需要给出数据，也需要说明数据为何可用。本作品把实际观测、标定区域和可靠尺度作为测量前提，并将不满足前提的结果保留为空。当前桥梁运行的世界坐标与速度留空，说明这一约束落实到了真实记录中。其设计价值是减少错误前提下的误解释，而不是用软件开关证明测量本身准确。','backend/processor.py','backend/data/runs/20261002-170236-c53990/vehicles.csv','plans/002-real-video-verification.json')
p('报告明确累计计数、短时流率、在场轨迹和车辆帧样本的不同含义，也排除预测行参与有效观测速样本统计。对本科生研究而言，这使数字能够与观察对象、时间范围和分母对应。已有功能证明系统能够保存这些区别，进一步评价需要检查实际计数和速度与独立真值的偏差。','backend/reports.py','backend/exporter.py','docs/features-analysis.md','docs/accuracy-evaluation.md')
p('当演示素材没有可靠道路尺寸时，系统仍保留车辆识别、轨迹和过线观察，把真实速度明确置为不可用。研究者因此可以继续检查不依赖米制尺度的部分，并知道测速实验还缺什么材料。把研究前提显示出来，有助于把一次演示整理为下一步可执行的实验，而不必用猜测数值补齐页面。','backend/processor.py','backend/reports.py','docs/accuracy-evaluation.md')

h('3.4 从报警结果延伸到证据与人工判断')
p('作品将报警定义为待核查的系统提示，并保留触发条件、截图和原视频。人工复核可以确认或否定单条提示，同时保持原有机器记录，使后续能够解释系统当时为何判断、人工又为何调整结论。这比仅保存一个报警数量更有利于发现规则和场景之间的问题。','backend/evidence.py','backend/reports.py','docs/features-analysis.md')
p('历史回放使用独立播放器，避免复查某次事件时打断当前分析。已有浏览器记录显示，历史播放器定位时当前分析仍继续完成，支持操作隔离的实现。该设计适合把演示延伸为研究资料检查，但尚未验证多人同时操作、长期保存和现场事件处置。','plans/002-verification.md','backend/evidence.py','backend/service.py')
p('例如研究疑似停车时，可以回放目标在提示前后的变化，检查它处于排队、暂时遮挡还是持续低速状态，再保存人工结论与理由。之后分析误报时，仍能看到当时的规则条件和机器事实。这样的操作使规则修改有可讨论的样例依据，而不是只凭即时界面印象决定参数。','backend/events/rules.py','backend/evidence.py','docs/features-analysis.md')

h('3.5 将重复使用与研究版本联系起来')
p('模板解决道路配置重复输入，版本绑定解决标注和评测条件变化后旧结果被误用。两者共同关注同一问题：结果只有与形成它的条件相连，才能被解释。作品在复用几何时保留尺度确认，在修改标注时撤销完整核验，给后续重复实验留下清楚的条件记录。','backend/camera_templates.py','backend/annotations.py','docs/camera-templates.md','docs/visual-evaluation.md')
p('这一特色的验证依托比例边界、尺度复制、版本冲突和保存重开等测试与界面操作。已有材料支持功能成立，尚不能证明配置时间缩短多少或不同隧道迁移效果如何。后续可在同机位重复配置与不同机位迁移中，记录操作步骤、错误项及重新标定需求，研究其实际使用收益。','tests/test_camera_templates.py','tests/test_annotations.py','plans/002-verification.md')

p('在同一机位重复分析时，保存的模板可以作为几何配置起点，使团队集中检查本轮真正改变的条件。标注版本则记录人工解释的变化，区分数据修改与模型变化。两项功能有助于组织重复研究，其具体收益可以通过后续任务记录评价；本阶段已有证据主要支持保存、复用和版本约束能够工作。','backend/camera_templates.py','backend/annotations.py','plans/002-verification.md')

h('4 科研实践与成果验证',1)
h('4.1 研究思路与本科生科研能力体现')
p('本作品的科研实践可以从问题界定、方案论证、原型建设和验证反思四个环节观察。问题界定把“监测交通”拆解为连续状态、道路位置、统计和复核；方案论证明确采用成熟模型、二维展示和本机录像范围；原型建设把分散模块连接为工作流程；验证反思则区分软件成立与实景精度。表4对应现有材料说明这些环节。','README.md','docs/features-analysis.md','docs/accuracy-evaluation.md','plans/002-verification.md')
table(4,'研究环节与可检查的成果材料',['研究环节','体现的研究能力','现有材料'],[
['问题界定','把应用需求转化为可研究的具体任务','需求文档、功能入口和作品实现范围'],
['方案选择','根据输入与条件选择并说明取舍','模型来源、素材抽查、二维与CPU方案'],
['系统建设','把感知、统计、展示和复核连接起来','前后端源码、历史报告与真实界面'],
['实验设计','区分软件测试、受控验证和真实评测','测试代码、隔离验收与人工评测工具'],
['分析反思','解释输出口径，识别结论适用边界','来源完整CSV、精度说明和待验证计划']],[3.0,6.4,7.6],['README.md','docs/accuracy-evaluation.md','plans/002-verification.md','backend/data/runs/20261002-170236-c53990/run.json'])
p('从选题能力看，作品将视频感知与交通观察联系起来，同时把完整系统限制在可以检查的范围内。隧道是应用目标，桥梁是当前真实验证场景；辨明这一区别，有助于提出下一阶段具体问题，而不是用作品名称替代实验依据。能够指出什么已经完成、什么尚未证明，是研究结论形成过程中的必要工作。','plans/002-real-video-verification.json','docs/accuracy-evaluation.md','plans/2026-10-02-implementation-status.md')
p('从方案论证看，采用预训练组件并不取消研究工作。需要分析目标尺寸和视角，检查现有模型的输出，决定哪些结果可以测量，以及如何使界面和数据含义一致。已有SSDLite抽查与默认模型调整提供了实践线索；仍需补充统一真值对照，才能把经验选型发展为更严格的实验结论。','README.md','models.json','docs/accuracy-evaluation.md')
p('从实验能力看，作品已经建设人工过线、事件标注和保存结果的工具，使误报、漏报与版本变化可以被记录。后续科研工作应利用这些工具提出明确问题、选择连续区间、独立建立真值，再分析错误来源。只有这样，测试工具才能成为研究方法的一部分，而不是以几个演示成功样例代替整体评价。','backend/annotations.py','backend/evaluation.py','docs/visual-evaluation.md')
h('4.2 实现条件与已有验证材料')
p('当前原型由Python/FastAPI后端和React/TypeScript前端实现，本次已记录环境为macOS 26.6 arm64、Python 3.13.0、PyTorch 2.11.0，采用CPU验证。前端使用Vite和ECharts组织界面与图表。环境信息说明已有运行条件，不代表不同平台、GPU或现场服务器已经完成测试。','output/competition-manual/evidence/environment.json','frontend/package.json','README.md')
p('验证材料分为三个层次。第一层是软件回归，检查几何、规则、同步和版本等约定；第二层是受控运行与浏览器操作，检查完整流程；第三层是保存的真实桥梁录像输出，检查实际素材是否可以经过系统。真实交通精度还需要独立真值，是下一层尚未完成的工作，不能从前三层通过情况推导。','tests/conftest.py','plans/002-verification.md','plans/002-real-video-verification.json','docs/accuracy-evaluation.md')
p('2026年10月3日的隔离复测记录为后端88项、前端27项通过，运行检查通过HTTP、可信来源空闲WebSocket和正常退出。这些结果支持当前软件流程和边界处理，不换算为检测准确率、测速精度或隧道验收。','output/competition-manual/evidence/validation-20261003.json','scripts/check_runtime.py')
p('已有浏览器验收覆盖模板保存与应用、草稿编辑与恢复、标注保存重开、报告、独立播放器和人工评测下载。它们证明操作链条在当时条件下能够完成。受控素材使用构造目标和标签，适合检验版本冲突、漏报记录等特定边界，却不能代表复杂交通场景的感知能力。','plans/002-verification.md','plans/002-browser-result.json')
p('后续浅色界面验收记录保存了桌面、平板、手机和窄屏的实际截图，前端31项测试及生产构建通过。这是2026年10月3日后续记录，不是本次重新运行结果。图6将实际历史报告与受控评测分为A、B两个面板；二者的材料与用途不同，不能组合成一次真实精度实验。','docs/linear-design/README.md','plans/features-report-desktop.jpg','plans/002-evaluation-result.jpg')
fig(6,'fig6-verification.png','功能验证界面  A实际桥梁历史报告  B构造标签人工评测',8.8)
h('4.3 真实桥梁录像的描述性结果')
p('选取统计段20261002-170236-c53990作为真实数据示例，是因为该段保存了视频内容身份、分析配置、道路场景、标定状态和CSV。来源为2160×3840、30FPS的桥梁俯拍录像，统计时间0.0—8.5秒，共256个统计帧。图7使用全部统计行，表5列出结果和各自口径；这些是描述性机器输出，没有独立人工真值。','backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','output/competition-manual/evidence/run-summary.json')
fig(7,'fig7-traffic.png','真实CSV曲线  累计机器过线与在场轨迹，全部256帧',6.2)
table(5,'真实桥梁运行结果及含义',['观察项目','实际记录','解释与限制'],[
['区间与统计帧','0.0—8.5s；256帧','末帧视频时间；原视频30FPS'],
['机器过线累计','19次','有限计数线、同段同ID去重；非人工计数'],
['轨迹数量','28个确认ID；最多同时13个','轨迹ID不等于物理车辆数；在场含预测'],
['车辆帧记录','1074观测＋1088预测＝2162行','同一车辆跨帧重复记录；不是2162辆车'],
['有效观测速样本','0/1074；覆盖率0%','未实测标定；速度与加速度不可用'],
['机器事件','0条','相关规则受标定门槛限制；不证明无异常']],[3.6,5.7,7.7],['backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/events.csv','output/competition-manual/evidence/run-summary.json'])
p('曲线能够回答机器计数怎样随视频时间增长、在场轨迹怎样变化，说明录像分析和记录保存已经连通。它不能回答机器漏掉了多少车辆，也不能把28个ID解释为真实车辆总数。累计过线19次应读取累计值的最大值，不能将逐帧累计列相加；正确理解口径是使用结果进行研究的前提。','backend/reports.py','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','docs/features-analysis.md')
p('该段末帧流率134.12辆/min由19÷8.5×60进行短时折算，并非连续一分钟的实际计数。既有CPU smoke记录的42.5FPS是处理吞吐，30FPS是原视频帧率，显示帧率还受推送和浏览器渲染影响，当前没有独立测量。相邻帧ID保留率98.9%只说明该次运行的连续性，不属于检测或跟踪精度指标。','docs/review-smoke-2026-10-02.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','docs/accuracy-evaluation.md')
p('未标定时车辆帧记录中的世界坐标、速度和加速度为空，机器事件为0。由此能够检查测量门槛确实生效，却不能推断录像中没有异常。这类没有输出的结果也应保留并解释，因为它反映了当前材料条件和系统使用前提。','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/events.csv','plans/002-real-video-verification.json')
h('4.4 人工评测流程的完成状态')
p('图6B来自受控流程验收。人工录入完整声明、过线与事件标签后，系统保存版本并计算一对一匹配结果，同时支持下载。截图中的计数误差0、事件TP2/FP1/FN1及66.7%的精确率和召回率来自刻意构造的标签，用于检查计算与界面一致性，不作为真实报警准确率。','plans/002-browser-result.json','plans/002-verification.md','backend/evaluation.py')
p('当前评测能够记录未映射到机器轨迹的漏报，避免只从机器报警列表中挑选确认事件而忽略完全未报的情况。标注修改会使旧结果过期，下载时检查版本。这些机制使下一阶段实验能够保存来源和条件，但软件不能自动证明人工已经核验了全部区间。','backend/annotations.py','backend/evaluation.py','docs/visual-evaluation.md')
p('当前没有可用于报告真实隧道检测准确率、真实事件表现或实测速度误差的完整真值材料。现有评测工具和计算路径尚未形成真实精度、对照或消融实验结论。当前成果应评价为功能原型及验证工具的完成，实际精度是后续研究问题。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','plans/2026-10-02-implementation-status.md')
h('4.5 进一步研究的问题与实验方案')
p('首先需要研究场景迁移：桥梁俯拍条件下可运行的感知与展示，在固定隧道入口和内部机位是否仍然适用。实验前应取得合法来源的连续录像，保存机位、时段、画面尺寸和相机稳定性等条件，并按明暗变化、遮挡和车流状况组织样本。具体区间和样本规模应在材料取得后确定。','docs/accuracy-evaluation.md','docs/camera-templates.md')
p('其次需要研究统计和事件误差。对同一计数线、同一区间独立记录人工过线；对三类事件说明定义与时间容差，完整核验后计算误报和漏报。应检查ID分裂、漏检、遮挡及方向配置等原因，而不只报告单个汇总指标。必要时由两名标注者独立核验并复查分歧，提高真值形成过程的可检查性。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','backend/evaluation.py')
p('测速研究需先取得可靠道路尺寸和独立参考速度，建立物理车辆、视频时间与系统ID的对应。在报告速度误差前，先说明多少真值样本成功匹配、多少没有有效输出，再给出MAE、RMSE等指标。道路尺寸和参考速度不能都从系统自己的输出中反推，否则不能构成独立验证[4]。','docs/accuracy-evaluation.md','https://arxiv.org/abs/1702.06441')
p('模型或规则对照应保持视频来源、区间、道路几何与真值相同，明确改变哪项条件。比较既要看错误，也要看处理耗时和不可用样本。现有条件对照和CSV可以支持记录，但不能替代独立实验。若新增检测评价，还需逐帧框和物理ID真值及相应导出，现有车辆明细不足以直接完成mAP或MOT类评价。','backend/reports.py','backend/exporter.py','docs/accuracy-evaluation.md')
p('在研究使用体验时，可以让实际使用者完成查找车辆、定位事件和核对报告等任务，记录操作问题与反馈。用户任务评价尚未开展，当前没有效率结果。与感知精度分别评价，能够说明二维展示和证据流程是否帮助使用者理解数据，避免用截图美观代替应用效果。','frontend/src/App.tsx','docs/features-analysis.md','plans/002-verification.md')

h('5 应用价值与推广前景',1)
h('5.1 当前可用的教学与离线研究场景')
p('现阶段较明确的使用方向是交通监测教学与离线录像研究。学习者可以从视频观察车辆，查看轨迹与道路关系，再通过计数线和时间窗口理解统计含义。预训练模型、标定和规则在同一原型中配合，使课程知识能够对应到实际输入和可保存结果。作品当前已有本机功能和运行材料支持这一使用方向。','README.md','docs/features-analysis.md','plans/002-verification.md')
p('离线研究可利用报告、CSV和独立回放，检查同一段录像的车辆变化与疑似事件。它适合用于整理样例、发现输出问题和准备人工标签。没有实测尺度时仍可研究检测、轨迹与机器过线，同时明确速度不可用；具备可靠尺度后再开展测速验证，能够避免为了展示完整而引入猜测数据。','backend/reports.py','backend/evidence.py','docs/accuracy-evaluation.md')
p('作为本科生科研原型，作品还为后续课题提供了可持续的基础：新问题可以从已有错误和材料条件中提出，新实验能够保存来源、配置和标签版本，研究结论可以回到原视频检查。当前价值在于形成可研究的工作环境，而不是宣称已经获得隧道生产系统的可靠性或经济收益。','backend/exporter.py','backend/annotations.py','docs/visual-evaluation.md')
p('教学中可以围绕一段连续录像设置研究任务：先独立记录过线，再查看机器统计，定位不一致的时刻，最后解释错误是否来自检测、轨迹或计数线配置。另一项任务是检查没有真实尺度时哪些量仍可观察、哪些量需要补充实测。系统现有回放、报告和标注工具能够支撑这些任务的设计，实际课程应用与学习效果仍待评价。','backend/reports.py','backend/evidence.py','backend/annotations.py','docs/accuracy-evaluation.md')

h('5.2 固定隧道机位的辅助监测潜力')
p('如果后续完成真实隧道录像、尺度与独立评测，系统可以进一步研究对固定机位的车道观察、断面统计和疑似停车、逆行等事件的辅助提示。二维视图用于理解位置关系，截图和原视频用于核查，历史报告用于复查区间变化。这些是由已有功能延伸出的应用方向，尚未形成现场部署案例。','docs/features-analysis.md','docs/camera-templates.md','plans/2026-10-02-implementation-status.md')
p('推广前需要确认机位几何、光照、车灯影响和遮挡等条件，并完成现场调试。PIARC将误报、漏报、入口眩光、照明变化和标定调试列为自动事件检测系统需要考虑的因素[1]。因此，本作品的桥梁模型选择和默认阈值不能直接当作隧道现场结论。','https://tunnelsmanual.piarc.org/en/equipment-systems-monitoringincident-detection/automatic-incident-detection','docs/accuracy-evaluation.md')
p('辅助监测的价值需要通过使用流程评价：系统能否给出可定位提示，人员能否回到原画面确认，漏报是否能在完整核验中被记录，以及配置变化是否可追溯。当前作品已经建设这些功能的原型，后续仍应验证稳定性和真实效果。系统没有完成自动执法或现场设备联动，不能用潜在应用代替已实现能力。','backend/evidence.py','backend/annotations.py','backend/service.py','plans/2026-10-02-implementation-status.md')
p('在固定机位试验阶段，可以先开展录像分析而不影响现场处置：按连续区间形成机器输出，由人工核验同一区间，再汇总常见错误和需要重新设置的条件。只有结果达到实际使用要求、运行链条通过验收后，才进一步研究持续接入。这样能够把应用目标转化为分阶段可检查的任务，也与当前离线原型形成衔接。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','backend/service.py')

h('5.3 分阶段推广与必要条件')
p('项目后续应先把本机录像分析做成可重复研究，再进入固定隧道机位评测，依据结果决定是否扩展实时流与多摄像头。图8展示阶段条件关系，表6列出方向与完成状态，后续阶段均有待取得材料和验收。','docs/accuracy-evaluation.md','plans/2026-10-02-implementation-status.md','backend/service.py')
fig(8,'fig8-roadmap.png','由研究原型走向应用的条件路线  后续阶段尚待验证',3.35)
table(6,'应用阶段与推广条件',['阶段','可研究或使用的方向','条件与完成状态'],[
['当前本机原型','教学、离线录像分析、证据复查','已有功能与运行材料；不报告真实交通精度'],
['真实隧道评测','固定机位车道统计与事件辅助提示','待补合法录像、实测尺度和完整独立标签'],
['实时流扩展','连续摄像头数据分析与异常恢复','RTSP底层接口存在；接入与断流验收未完成'],
['多摄像头与现场使用','扩大覆盖、统一历史与运行管理','并发隔离、容量、权限和长期运行均待建设或验证']],[3.5,6.0,7.5],['backend/vision/video_reader.py','backend/service.py','docs/accuracy-evaluation.md','plans/2026-10-02-implementation-status.md'])
p('实时流扩展需要同时检查视频时间、连接失败、恢复行为和数据保存，而不仅是让底层接口接收地址。多摄像头需要独立会话、资源管理和跨来源记录，当前单会话架构不能直接宣称支持。长期运行还需考虑存储容量、来源缺失、磁盘异常和服务恢复，已有短时smoke没有覆盖这些条件。','backend/vision/video_reader.py','backend/service.py','scripts/check_runtime.py','plans/2026-10-02-implementation-status.md')
p('现场使用还需核对视频和权重的使用条件，明确数据访问、保留和人工处置责任。素材与模型授权、合作条件和使用责任均需在实际推广前核实，当前没有落地或经济收益记录。最终是否投入现场，应依据真实评测和运行验收决定；当前适合继续作为可解释、可复查的研究原型完善。','models.json','backend/service.py','plans/2026-10-02-implementation-status.md')
h('5.4 作品的持续研究价值')
p('项目的后续价值取决于能否把当前原型转化为重复研究的基础。首先利用真实材料检验现有假设，再根据误差分析调整模型、几何或规则，最后在相同条件下复测。这样的路线能够让改进与证据对应，也便于团队在后续课程、毕业设计或科研训练中持续积累资料。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','backend/reports.py')
p('即使某种场景下表现不足，系统仍可保存问题区间、对应配置和人工标注，帮助解释限制来自目标识别、轨迹关联、道路测量还是规则定义。失败和不可用样本应进入研究记录，而不是删除后仅展示效果较好的片段。作品中的来源保存与版本评测功能，为这种持续分析提供了工具。','backend/exporter.py','backend/annotations.py','docs/accuracy-evaluation.md')

h('6 结论与后续工作',1)
p('本作品围绕交通视频的状态理解、统计与核查需求，完成了本机视频感知与二维数字孪生原型，连接车辆识别、道路配置、同帧展示、报告、证据和人工评测。已有软件回归、浏览器操作和真实桥梁运行记录支持当前功能。作品的主要特色是信息之间具有明确对应关系，测量前提受到约束，结果能够回到来源和条件检查。','docs/features-analysis.md','docs/visual-evaluation.md','plans/002-verification.md','plans/002-real-video-verification.json')
p('从本科生科研实践的角度，作品提供了可检查的需求分析、方案选择、系统建设和验证反思材料。下一阶段应以真实隧道素材和独立真值为重点，研究场景适应性、计数与事件错误以及速度测量条件，再根据结果决定实时流和现场扩展。当前不以原型完成代替精度验证，也不以应用目标代替部署成果；后续价值需要在连续的实测与复查中建立。','docs/accuracy-evaluation.md','docs/camera-templates.md','plans/2026-10-02-implementation-status.md')
h('参考文献',1)
refs=[
('[2] Veres M, Moussa M. Deep Learning for Intelligent Transportation Systems: A Survey of Emerging Trends[J]. IEEE Transactions on Intelligent Transportation Systems, 2020, 21(8): 3152-3168. DOI: 10.1109/TITS.2019.2929020.','https://doi.org/10.1109/TITS.2019.2929020'),
('[3] Sochor J, Juránek R, Herout A. Traffic Surveillance Camera Calibration by 3D Model Bounding Box Alignment for Accurate Vehicle Speed Measurement[J]. Computer Vision and Image Understanding, 2017, 161: 87-98. DOI: 10.1016/j.cviu.2017.05.015.','https://doi.org/10.1016/j.cviu.2017.05.015'),
('[4] Sochor J, Juránek R, Špaňhel J, et al. Comprehensive Data Set for Automatic Single Camera Visual Speed Measurement[J]. IEEE Transactions on Intelligent Transportation Systems, 2019, 20(5): 1633-1643. DOI: 10.1109/TITS.2018.2825609.','https://doi.org/10.1109/TITS.2018.2825609'),
('[5] dronefreak. visdrone-yolo11n 模型卡[EB/OL]. [2026-10-03].','https://huggingface.co/dronefreak/visdrone-yolo11n'),
('[6] Ultralytics. Ultralytics YOLO11 官方文档[EB/OL]. [2026-10-03].','https://docs.ultralytics.com/models/yolo11/'),
('[7] Torchvision. ssdlite320_mobilenet_v3_large 官方文档[EB/OL]. [2026-10-03].','https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.detection.ssdlite320_mobilenet_v3_large.html'),
('[8] Zhu P, Wen L, Bian X, et al. Vision Meets Drones: A Challenge[EB/OL]. 2018. arXiv:1804.07437.','https://arxiv.org/abs/1804.07437'),
('[9] Bewley A, Ge Z, Ott L, et al. Simple Online and Realtime Tracking[C]. ICIP, 2016. DOI: 10.1109/ICIP.2016.7533003.','https://doi.org/10.1109/ICIP.2016.7533003'),
('[10] OpenCV. Geometric Image Transformations[EB/OL]. [2026-10-03].','https://docs.opencv.org/4.13.0/da/d54/group__imgproc__transform.html'),
]
refs.append(('[1] PIARC. Road Tunnels Manual: Video Automatic Incident Detection[EB/OL]. [2026-10-03].','https://tunnelsmanual.piarc.org/en/equipment-systems-monitoringincident-detection/automatic-incident-detection'))
refs.sort(key=lambda x:int(re.match(r'\[(\d+)\]',x[0])[1]))
from docx.opc.constants import RELATIONSHIP_TYPE as RT
for text,url in refs:
    pa=doc.add_paragraph(text);pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(1)
    for r in pa.runs:r.font.size=Pt(10.5)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz);rr.append(rp)
    tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=' 在线来源';rr.append(tt);link.append(rr);pa._p.append(link)



path=OUT/'基于视频感知的隧道交通数字孪生系统作品说明书.docx';doc.save(path)
(OUT/'evidence/manuscript.md').write_text('\n\n'.join(draft),encoding='utf-8')
with (OUT/'evidence/claims.csv').open('w',newline='',encoding='utf-8-sig') as f:
    wr=csv.DictWriter(f,fieldnames=list(claims[0]));wr.writeheader();wr.writerows(claims)
(OUT/'evidence/displays.json').write_text(json.dumps(displays,ensure_ascii=False,indent=2))
print(path)
