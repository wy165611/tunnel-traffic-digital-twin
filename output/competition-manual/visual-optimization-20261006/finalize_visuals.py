from pathlib import Path
import json,csv,re,hashlib,zipfile
from docx import Document
from lxml import etree
from pypdf import PdfReader

OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[2];E=OUT/'evidence';OLD=OUT.parent/'revision-undergraduate'
RENDER=Path('/private/tmp/twin-visual-qa4')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write(name,value):(E/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
manifest=json.loads((E/'source_manifest.json').read_text());by={s['path']:s['evidence_id'] for s in manifest['sources']}
baseline=json.loads((E/'input-baseline.json').read_text())
drift=[]
for s in manifest['sources']:
    if not s.get('sha256'):continue
    p=Path(s['path']);p=p if p.is_absolute() else ROOT/p
    if not p.is_file() or sha(p)!=s['sha256']:drift.append(s['path'])
assert not drift,drift
assert sha(OLD/'基于视频感知的隧道交通数字孪生系统作品说明书.docx')==baseline['previous_docx_sha256']

rows=list(csv.DictReader((E/'claims.csv').open(encoding='utf-8-sig')));claims={r['claim_id']:r for r in rows};known={s['evidence_id'] for s in manifest['sources']}
claim_audit=[]
for line in (E/'manuscript.md').read_text().splitlines():
    hit=re.search(r'\[claim:(C\d+)\]',line)
    if not hit:continue
    content=re.sub(r'\[(?:claim|evidence):[^\]]*\]','',line)
    r=claims[hit[1]];assert r['claim_text_sha256']==hashlib.sha256(' '.join(content.split()).encode()).hexdigest()
    assert all(v in known for v in r['evidence_ids'].split(',') if v)
    claim_audit.append({'claim_id':hit[1],'text_hash_binding':'PASS','evidence_binding':'PASS','human_verification':'pending'})
write('claim-binding-audit.json',claim_audit)

figure_sources={
1:['docs/features-analysis.md','docs/visual-evaluation.md'],
2:['backend/processor.py','backend/service.py','frontend/src/App.tsx'],
3:['docs/linear-design/desktop-full.jpg','docs/linear-design/README.md','frontend/src/components/DigitalTwin.tsx'],
4:['backend/events/rules.py','backend/evidence.py','docs/features-analysis.md'],
5:['frontend/src/utils/snapshotState.ts','backend/processor.py','backend/evidence.py','backend/annotations.py'],
6:['plans/features-report-desktop.jpg','plans/002-evaluation-result.jpg','plans/002-browser-result.json'],
7:['backend/data/runs/20261002-170236-c53990/traffic_statistics.csv'],
8:['docs/accuracy-evaluation.md','plans/2026-10-02-implementation-status.md','backend/vision/video_reader.py']}
processing={1:'Project-need schematic, not an experimental finding',2:'Source-derived local single-session topology, simplified for the undergraduate audience',3:'Literal crops and proportional enlargement from one actual desktop-full.jpg snapshot; overview crops [389,542,700,990] and [1015,542,1235,990]; ID 22 inset crops [521,561,668,639] and [1090,605,1237,683]; labels added outside screenshots; no scene data altered',4:'Source-derived event evidence workflow with observation and scale prerequisites',5:'Qualitative synthesis of implemented design relationships; no benefit percentages claimed',6:'Panel A literal report crop [180,249,797,355] from features-report-desktop.jpg, run 20261002-165026-852c78; panel B literal evaluation crop [325,548,1000,697] from 002-evaluation-result.jpg with constructed labels; distinct sessions and evidence roles',7:'Matplotlib step plots of all 256 CSV records, timestamp/flow/vehicle_count; distinct axes and units; no smoothing or interpolation; not independent traffic truth',8:'Conditional application roadmap; current prototype and future unvalidated stages explicitly distinguished'}
ds=json.loads((E/'displays.json').read_text())
for d in ds:
    if d['kind']=='figure':
        n=d['number'];d.update({'evidence':[by[s] for s in figure_sources[n]],'asset_sha256':sha(OUT/'assets'/d['asset']),'processing':processing[n],'human_verification':'pending','material_rights':'User-supplied project materials; competition redistribution rights to be confirmed by team'})
write('displays.json',ds)

run=ROOT/'backend/data/runs/20261002-170236-c53990'
stats=list(csv.DictReader((run/'traffic_statistics.csv').open(encoding='utf-8-sig')))
vehicles=list(csv.DictReader((run/'vehicles.csv').open(encoding='utf-8-sig')))
events=list(csv.DictReader((run/'events.csv').open(encoding='utf-8-sig')))
assert len(stats)==256 and max(float(r['timestamp']) for r in stats)==8.5
assert max(int(r['flow']) for r in stats)==19 and max(int(r['vehicle_count']) for r in stats)==13
assert len(vehicles)==2162 and sum(r['observed']=='True' for r in vehicles)==1074
assert sum(r['observed']=='False' for r in vehicles)==1088 and len({r['vehicle_id'] for r in vehicles})==28
assert not events and all(not r[k] for r in vehicles for k in ['x','y','speed','acceleration'])
summary=json.loads((E/'run-summary.json').read_text());assert sha(ROOT/'backend/data/videos/demo.mp4')==summary['source_sha256']

docx=OUT/'基于视频感知的隧道交通数字孪生系统作品说明书.docx';doc=Document(docx)
with zipfile.ZipFile(docx) as z:
    xml=etree.fromstring(z.read('word/document.xml'));settings=etree.fromstring(z.read('word/settings.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
assert len(doc.tables)==6 and len(doc.inline_shapes)==8
assert not settings.xpath('//w:documentProtection',namespaces=ns)
assert not xml.xpath('//m:oMath',namespaces=ns)
cites=xml.xpath('//w:r[w:rPr/w:vertAlign[@w:val="superscript"]]/w:t/text()',namespaces=ns)
first=[]
for s in cites:
    hit=re.search(r'\[(\d+)(?:-(\d+))?\]',s);assert hit,s
    for n in range(int(hit[1]),int(hit[2] or hit[1])+1):
        if n not in first:first.append(n)
assert first==list(range(1,11)),first
figs=[int(re.match(r'图(\d+)',p.text)[1]) for p in doc.paragraphs if re.match(r'图\d+  ',p.text)]
tables=[int(re.match(r'表(\d+)',p.text)[1]) for p in doc.paragraphs if re.match(r'表\d+  ',p.text)]
assert figs==list(range(1,9)) and tables==list(range(1,7))
refs=json.loads((E/'references.json').read_text());assert len(refs)==10 and sum(r['publication_type']=='journal' for r in refs)==3
for table in doc.tables:
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    for cell in table.rows[0].cells:
        assert cell._tc.xpath('./w:tcPr/w:shd/@w:fill')==['183E59']
        assert all(r.font.color.rgb is not None and str(r.font.color.rgb)=='FFFFFF' for p in cell.paragraphs for r in p.runs)
    assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)
pdf=next(RENDER.glob('*.pdf'));pages=len(PdfReader(pdf).pages);assert 14<=pages<=20
assert len(list(RENDER.glob('page-*.png')))==pages
write('delivery-audit.json',{'date':'2026-10-06','docx_sha256':sha(docx),'pages':pages,'figures':8,'native_editable_tables':6,'reference_count':10,'journal_papers':3,'superscript_citation_runs':len(cites),'citations_ordered_by_first_use':True,'claim_bindings':'PASS','numeric_consistency':'PASS','current_input_hash_mismatches':drift,'preexisting_source_changes_since_previous_revision':baseline['preexisting_changes_since_prior_revision'],'previous_docx_preserved':True,'source_video_identity_rechecked':True,'new_experiments':False,'render_directory':str(RENDER),'visual_review':'pending final full-page inspection','human_verification':'pending','file_size_bytes':docx.stat().st_size})
coverage=json.loads((E/'reporting_coverage.json').read_text());coverage['topics']['software_validation']='4.2: historical 88 backend/27 frontend tests, later historical 31 frontend tests';coverage['topics']['display_provenance']='Figures 1–8 plus displays.json';coverage['submission_ready']=False;write('reporting_coverage.json',coverage)
cm=json.loads((E/'consistency_manifest.json').read_text());cm['numeric_facts'].append({'fact_id':'N_VISUAL_001','concept':'later_frontend_validation','section':'4.2','value':31,'unit':'test','analysis_set':'historical_UI_acceptance_20261003','evidence_ids':[by['docs/linear-design/README.md']]});write('consistency_manifest.json',cm)
write('revision-record.json',{'date':'2026-10-06','request':'Optimize illustrations, tables, printing clarity and narrative distribution while preserving undergraduate research emphasis','previous_docx_sha256':baseline['previous_docx_sha256'],'changes':['New background need diagram','Compact source-derived architecture','Actual light-theme paired screenshot with ID 22 magnification','Clear evidence workflow','Design relationship diagram','Actual report and controlled evaluation separated by labelled panels','All real CSV rows plotted with separate statistical grains and units','Conditional application roadmap','Dark blue native table headers, alternating rows, 11 pt text and larger cell padding','Post-table spacing and protected paragraph/page relationships'],'experiments_added':False,'project_source_or_data_modified':False,'preexisting_drift_count':len(baseline['preexisting_changes_since_prior_revision']),'human_approval':'pending'})

check=f'''图文优化说明与证据来源及待补材料
2026年10月6日

一、本次交付
说明书共{pages}页、8幅图、6张原生Word表格；正文与表格可编辑，正文引用为数字上标。保留本科生科研实践重点，新增需求关系、特色联系和条件路线图，重绘架构与真实CSV曲线；真实同帧截图增加ID 22局部放大，实际报告与构造标签评测分别标明。统一图表编号、单位、表格配色和留白。已保留旧稿。

二、主要证据
图1、2、4、5、8依据项目需求和源码绘制，是关系或流程示意。图3来自docs/linear-design/desktop-full.jpg的真实浅色界面截图，裁切与放大同一帧，未标定；图6A来自plans/features-report-desktop.jpg，图6B来自plans/002-evaluation-result.jpg，B为受控构造标签验收。图7使用统计段20261002-170236-c53990全部256行CSV，累计机器过线19次，在场轨迹最多13个，包含预测；1074观测＋1088预测＝2162条车辆帧记录。没有独立真值，真实速度不可用。
验证沿用既有记录：2026年10月3日较早隔离复测88项后端、27项前端通过；后续浅色界面验收记录31项前端通过及构建通过。本次仅核对材料、制图、编辑与排版，没有新增系统实验。6个源码/文档文件相较旧稿已有界面更新，本次按现有文件建立新基线；基线与交付后核对一致。原稿、源码、视频、模型、标定及历史分析数据保持完整。
主文10条真实参考来源沿用已核对记录，含IEEE Transactions on Intelligent Transportation Systems和Computer Vision and Image Understanding的3篇期刊论文，并保留模型、方法与PIARC原始来源。图表来源、裁切位置、哈希、文字证据和一致性记录保存在本目录evidence/。

三、待补材料
团队、学校、指导教师及真实个人贡献；最终赛道和正式模板；合法真实隧道录像；固定机位与道路尺寸实测；完整独立过线和事件标签、独立参考测速；素材和模型使用条件；实时流、长时运行、并发及现场验收。当前是桥梁素材验证、二维展示和本机单分析会话，不能推导真实隧道准确率或测速误差。继续按用户20页内要求制作，正式提交前按确定的竞赛模板核对。

四、编制工具与方法
使用Codex、scientific-writing与Documents技能辅助编辑和排版。方法工具引用单独保留，不混入交通技术文献：
Kassis T, Agarwal V, He Y, Patel D, Brueckner A M. Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents[EB/OL]. 2026. arXiv:2609.00065. DOI:10.48550/arXiv.2609.00065. https://arxiv.org/abs/2609.00065 。原始记录已核对，该预印本不作为本项目实验依据。团队最终事实、贡献及提交确认仍由人工完成，证据记录不冒充人工签署。
'''
(OUT/'图文优化说明与证据来源及待补材料清单.txt').write_text(check,encoding='utf-8')
print(f'Structural and numeric audit PASS: {pages} pages, 8 figures, 6 native tables, {len(cites)} superscript citations, preserved inputs.')
