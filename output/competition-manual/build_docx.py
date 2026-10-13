from pathlib import Path
import csv,json,hashlib,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.5)
sec.left_margin=sec.right_margin=Cm(2)
sec.footer_distance=Cm(1.25)
for name,size in [('Normal',12),('Title',17),('Heading 1',12),('Heading 2',12),('Caption',10.5)]:
    s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'STSong')
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        s._element.rPr.rFonts.attrib.pop(qn('w:'+key),None)
    if s._element.pPr is not None:
        for border in list(s._element.pPr.findall(qn('w:pBdr'))):s._element.pPr.remove(border)
    pf=s.paragraph_format; pf.line_spacing=1; pf.space_before=Pt(0); pf.space_after=Pt(1)
    if name.startswith('Heading'): pf.space_before=Pt(3); pf.space_after=Pt(1); s.font.bold=True
    if name=='Normal': pf.first_line_indent=Pt(24)
    else: pf.first_line_indent=Pt(0)
    if name=='Title':s.font.bold=True
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.core_properties.title='基于视频感知的隧道交通数字孪生系统作品说明书'
doc.core_properties.author=''; doc.core_properties.subject='智能交通竞赛作品说明书'
doc.core_properties.keywords='车辆检测 多目标跟踪 道路标定 二维数字孪生 人工评测'
doc.core_properties.comments=''
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
r=footer.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)

manifest=json.loads((OUT/'evidence/source_manifest.json').read_text())
E={x['path']:x['evidence_id'] for x in manifest['sources']}
claims=[]; draft=[]; displays=[]
def evidence(*paths): return [E[p] for p in paths]
def record(text,paths):
    ids=evidence(*paths) if paths else []
    c=f'C{len(claims)+1:03d}'
    claims.append({'claim_id':c,'claim_text_sha256':hashlib.sha256(' '.join(text.split()).encode()).hexdigest(),'evidence_ids':','.join(ids),'verification_status':'human_pending','assistant_checked_on':'2026-10-03'})
    draft.append(f'{text} [claim:{c}] [evidence:{",".join(ids)}]')
def add_cited_text(pa,text):
    for part in re.split(r'(\[\d+(?:[-,]\d+)*\])',text):
        run=pa.add_run(part)
        if re.fullmatch(r'\[\d+(?:[-,]\d+)*\]',part):
            run.font.superscript=True; run.font.size=Pt(9)
def p(text,*paths,small=False):
    pa=doc.add_paragraph()
    add_cited_text(pa,text)
    if small:
        pa.paragraph_format.first_line_indent=Pt(0)
        for r in pa.runs:r.font.size=Pt(10.5)
    record(text,paths)
    if text.startswith('H为3×3'):pa.paragraph_format.keep_with_next=True
    return pa
def h(text,level=2):
    doc.add_paragraph(text,'Heading '+str(level));draft.append('## '+text)
def newpage():pass
def fig(n,name,caption,height):
    im=OUT/'assets'/name
    from PIL import Image
    w,hh=Image.open(im).size; width=min(17,height*w/hh)
    pa=doc.add_paragraph(); pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
    pa.alignment=WD_ALIGN_PARAGRAPH.CENTER; pa.paragraph_format.keep_with_next=True
    pic=pa.add_run().add_picture(str(im),width=Cm(width)); pic._inline.docPr.set('descr',caption)
    cp=doc.add_paragraph(f'图{n}  {caption}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    displays.append({'number':n,'kind':'figure','asset':name,'caption':caption,'processing':'local diagram or screenshot crop; no fact alteration'})
def table(n,title,headers,rows,widths,paths):
    cp=doc.add_paragraph(f'表{n}  {title}','Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for col,w in zip(tb.columns,widths):col.width=Cm(w)
    for cell,w in zip(tb.rows[0].cells,widths):cell.width=Cm(w)
    for cell,text in zip(tb.rows[0].cells,headers):cell.text=text
    for row in rows:
        cells=tb.add_row().cells
        for cell,w,text in zip(cells,widths,row):cell.width=Cm(w);cell.text=str(text)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for tag in ['top','bottom','left','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF3');tcpr.append(sh)
            for pa in cell.paragraphs:
                pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0);pa.paragraph_format.line_spacing=1
                if ci==0:pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if n==1 and ri<len(tb.rows)-1:pa.paragraph_format.keep_with_next=True
                raw=pa.text
                pa.clear();add_cited_text(pa,raw)
                for r in pa.runs:
                    if not r.font.superscript:r.font.size=Pt(10.5)
                    r.bold=ri==0
    record('；'.join([title]+[' / '.join(str(x) for x in row) for row in [headers]+rows]),paths)
    displays.append({'number':n,'kind':'table','title':title,'native_word_table':True,'evidence':evidence(*paths)})
def eq(prefix,numerator,denominator,suffix):
    pa=doc.add_paragraph();pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0)
    m=OxmlElement('m:oMath')
    def mr(t):
        r=OxmlElement('m:r'); x=OxmlElement('m:t');x.text=t;r.append(x);return r
    m.append(mr(prefix));f=OxmlElement('m:f')
    for tag,t in [('num',numerator),('den',denominator)]:
        z=OxmlElement('m:'+tag);z.append(mr(t));f.append(z)
    m.append(f);m.append(mr(suffix));pa._p.append(m)

# Page 1: front matter and scope, matching the official no-cover requirement.
pa=doc.add_paragraph('基于视频感知的隧道交通数字孪生系统','Title');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['参赛团队：待补充    指导教师：待补充','学校及院系：待补充    所在城市及邮编：待补充']:
    pa=doc.add_paragraph(text);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(0)
pa=doc.add_paragraph('作品内容简介','Heading 1');pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('本作品面向隧道交通监测中视频信息难以转化为可解释状态、异常报警缺少复核依据等问题，构建本机视频感知与二维道路数字孪生系统。系统采用现成预训练车辆检测模型，以SORT风格的Kalman预测和Hungarian匹配维持轨迹ID；用轨迹框底边中心定位车辆，在可靠实测尺度下通过四点Homography映射至米制平面，并利用实际视频时间戳和窗口最小二乘拟合计算速度。系统实现有限计数线过线统计、车道和类别分析、10/30/60秒车道窗口，以及停车、急减速和逆行规则。FastAPI与WebSocket发布同帧图像和状态，React展示视频、位置、朝向、轨迹、统计与事件；统计段归档、报警截图、独立回放、机位模板和版本化人工评测支持结果追溯。已有桥梁俯拍素材的运行记录与软件回归支撑功能实现，尚无真实隧道真值及实测测速精度结果。作品未训练或改进检测网络，当前为二维展示和本机单分析会话。未实测标定时，米制位置、速度、加速度与相关报警保持不可用。','backend/processor.py','backend/service.py','backend/vision/detector.py','docs/features-analysis.md','docs/visual-evaluation.md','plans/002-verification.md')
p('关键词：视频车辆检测；多目标跟踪；道路标定；二维数字孪生；证据复核',small=True)
h('1 研究背景',1)
h('1.1 隧道交通监测需求与应用问题')
p('隧道监测需将视频中的车辆在场、过线和疑似异常联系到道路空间。缺少尺度、时间和原画面，测速与报警难以解释；仅复核已报事件也无法发现完全漏报，因此系统同时提供独立全区间核验。','docs/features-analysis.md','docs/accuracy-evaluation.md')
h('1.2 相关研究与技术现状')
p('交通深度学习综述强调时空及场景差异[1]；单目测速须标定尺度[2]和独立真值[3]。本作品采用实测四点标定，未采用文献[2]的三维模型自动尺度推断。YOLO11和SSDLite提供预训练检测组件[4-6]；VisDrone研究聚焦无人机检测与跟踪[7]，为俯拍选型提供依据。SORT结合Kalman滤波与Hungarian关联[8]。本作品采用已有方法，重点实现测量门控、同帧展示和可追溯评测。','models.json','backend/vision/detector.py','backend/vision/tracker.py')
h('1.3 本作品的目标及实现范围')
p('当前完成本机录像分析与复核，主要素材为桥梁俯拍，隧道场景待验证。孪生为二维状态映射。底层读取保留RTSP接口，现有入口以文件为主，实时流未验收。','backend/vision/video_reader.py','backend/service.py','plans/002-verification.md')
h('2 设计原理与关键技术',1)
h('2.1 系统总体架构与数据流程')
p('图1为单分析会话的源码架构，表1列出模块分工。','backend/service.py','backend/processor.py','backend/main.py')
fig(1,'fig1-architecture.png','系统总体架构  依据当前源码绘制',3.0)
table(1,'系统模块及功能',['模块','主要功能','输出与约束'],[
['视频与感知','解码PTS、检测适配、ID跟踪','区分实际观测和预测'],
['道路与交通','四点标定、测速、过线及窗口统计','米制测量须可靠实测尺度'],
['展示与证据','双视图同步、事件截图及独立回放','二维、本机单分析会话'],
['报告与评测','CSV/HTML、模板、人工标注与版本评测','机器结果、复核与真值分别存储']],[2.8,7.1,7.1],['backend/processor.py','backend/reports.py','backend/evidence.py','backend/annotations.py'])

newpage()
h('2.2 视频车辆检测与多目标跟踪')
p('预训练权重见表2，加载前按models.json校验SHA256。类别统一为car、motorcycle、bus、truck，VisDrone的van并入car。桥梁初期抽查SSDLite漏检较多，默认选用VisDrone YOLO11n；尚无统一真值的模型精度对比。','backend/vision/detector.py','backend/vision/weights.py','models.json','README.md')
table(2,'预训练模型来源与选型依据',['模型','来源与权重','场景与选型依据'],[
['VisDrone YOLO11n','dronefreak的VisDrone微调权重[4]','俯拍小目标演示默认；由发布者训练'],
['COCO YOLO11n','Ultralytics官方yolo11n.pt[5]','通用道路视角备选；需新素材复核'],
['SSDLite MobileNetV3','Torchvision官方COCO权重[6]','320×320轻量推理；当前俯拍抽查漏检']],[4,6.2,6.8],['models.json','backend/vision/detector.py','README.md'])
p('默认每隔2帧检测，其间预测（图2）。框中心、宽高及变化率组成8维状态，观测为4维；IoU门控后以1−IoU代价进行Hungarian一对一关联。满足最少命中数才输出，超龄清理。没有外观重识别，遮挡和ID分裂可能导致重复或漏计。','backend/vision/tracker.py','config.yaml')
fig(2,'fig2-pipeline.png','视频感知与交通分析流程  观测门控区分检测与预测',2.7)
p('observed=true为检测关联后的轨迹，false为预测；预测用于展示，不更新测速、计数或规则。跟踪保留30帧，展示另有限制漏检时长，避免长期预测框输出。','backend/processor.py','backend/vision/tracker.py')
h('2.3 道路标定与交通参数计算')
p('将轨迹框还原至原图，取底边中心近似道路接触点。依次选左上、右上、右下、左下角，并输入对应矩形的实测宽W、长L，求解H[9]。仅有效区域内的点参与米制测量；图3为概念示意。','backend/processor.py','backend/calibration/homography.py')
fig(3,'fig3-calibration.png','四点道路标定与坐标映射  概念示意，非实测道路',3.7)
eq('X = ','h₁₁u + h₁₂v + h₁₃','h₃₁u + h₃₂v + h₃₃','； Y = (h₂₁u + h₂₂v + h₂₃) / (h₃₁u + h₃₂v + h₃₃)')
p('H为3×3单应矩阵，u、v以px计，X、Y以m计。对实际观测的米制点，令时间偏差Δtᵢ=tᵢ−平均时间，以窗口最小二乘拟合二维速度：','backend/calibration/homography.py','backend/traffic/speed_estimator.py')
eq('v = ','Σ Δtᵢ (rᵢ − r̄)','Σ (Δtᵢ)²','； s = 3.6 × mean₃(‖v‖)')
p('rᵢ为道路点，v单位m/s，s单位km/h。默认10个观测点，至少5点且跨度≥0.1s才测速，最近3次模长取均值；加速度由相邻平滑速度差计算。采用解码PTS，断档过长或离开有效区域时重置缓存。平面假设、畸变、相机移动和接触点偏差均待实测验证。','backend/traffic/speed_estimator.py','backend/processor.py','backend/vision/video_reader.py')
p('过线判定检查连续实际观测点的连线与有限计数线相交，同一ID在一段内只计一次；车道以接触点落入人工多边形确定，类别和车道过线统计取穿越时归属。累计值不逐帧求和。10/30/60s窗口取(t−w,t]，覆盖为min(w,t−段起点)，不足时标记不完整；每分钟流率在不足60s时为短时折算。','backend/traffic/flow_counter.py','backend/processor.py','docs/features-analysis.md')

newpage()
h('2.4 视频与二维孪生同步')
p('后端整帧完成后发布帧号、视频时间、JPEG、车辆、轨迹、统计和事件，WebSocket按修订号推送。前端以会话与修订号过滤旧响应，解码图像与来源快照成对提交，双视图共用状态，防止迟到图像错配位置。图4为未标定横向桥梁的实际同帧界面。','backend/service.py','backend/main.py','frontend/src/utils/snapshotState.ts','frontend/src/App.tsx')
fig(4,'fig4-twin.png','桥梁视频与二维孪生同帧对照  实际界面，未标定',4.5)
p('未标定时保留原图比例，标定后以米尺度等比例展示，按角点调整方向和坐标轴。朝向由近期观测位移估计，静止时不等于真实车头方向。实时全局均速可包含预测轨迹保留的最近速度；报告与车道窗口只用实际观测，二者口径不同。','frontend/src/utils/twinGeometry.ts','frontend/src/components/DigitalTwin.tsx','backend/processor.py','backend/traffic/traffic_metrics.py','backend/reports.py')
h('2.5 异常事件与证据复核')
p('异常停车、急减速和逆行阈值见表3，仅处理实际观测；无有效速度时清缓存并返回无报警，均依赖标定和足够观测。停车不含停车区语义，急减速受测速抖动影响；逆行需正确车道方向，ID变化可能中断判断。','backend/events/rules.py','backend/processor.py')
table(3,'关键参数单位与启用条件',['参数','默认值','条件或统计含义'],[
['检测与匹配','置信度0.40；间隔2帧；IoU 0.20','输入960；最少命中2次'],
['测速窗口','10点；至少5点；跨度≥0.1s','实测四点尺度；输出km/h、m/s²'],
['异常停车','速度<3km/h，持续≥3s','有效速度与实际观测'],
['急减速','加速度<−4m/s²','规则使用≥0.5s速度差'],
['逆行','余弦<−0.6，持续≥1s','速度>3km/h；合法方向已配置'],
['冷却与车道窗口','同ID同事件5s；10/30/60s','视频时间；窗口不足标记不完整']],[3.2,6.3,7.5],['config.yaml','backend/events/rules.py','backend/traffic/speed_estimator.py','backend/traffic/flow_counter.py'])
p('如图5，报警触发时原子保存稳定事件ID、观测目标截图和不可变事实，记录视频身份、时间、规则、轨迹与标定。独立播放器先核验原视频内容身份，再通过Range回放事件邻域；来源缺失或改变时禁止回放，截图仍保留。人工确认、误报及备注另存，不修改机器事件或实时会话。','backend/evidence.py','docs/features-analysis.md')
fig(5,'fig5-evidence.png','异常事件证据与人工复核流程  依据源码绘制',2.1)

newpage()
h('2.6 历史报告 机位模板和人工评测')
p('统计段保存来源、参数、道路与标定快照。暂停继续保留原段，定位或换条件形成新段，旧CSV保留。报告流式读取CSV，累计过线取最大值；均速以有效观测车辆帧为样本，覆盖率为有效测速样本数/实际观测车辆帧数，不是每车均速或检测准确率。图6为桥梁报告区间与累计值，未标定速度留空。','backend/service.py','backend/exporter.py','backend/reports.py','docs/features-analysis.md')
fig(6,'fig6-report.png','桥梁报告指标区  实际界面，19次机器过线',3.0)
p('CSV及离线HTML/SVG可导出。对照要求同视频身份、区间差≤0.1s、几何与尺度相同，另列模型和规则差异；数量变化不表示精度提升。趋势图降采样，原CSV保留明细。','backend/reports.py','docs/features-analysis.md')
p('固定机位模板保存已保存的几何和可选实测尺度，按版本管理。同比例分辨率按像素宽高比映射，宽高比偏差超过1%仅允许预览；默认应用几何会清除真实尺度。即使图像比例相同，也必须人工确认机位、视野和实测尺寸仍有效，才能单独复制尺度。裁切、旋转、移动或变焦后须重新核验，模板不能代替标定。','backend/camera_templates.py','docs/camera-templates.md')
p('人工核验已关闭段的完整录像，记录物理车辆过线、事件与已核验类别。标注绑定运行、来源、配置指纹、尺寸和revision，修改撤销完整声明、使旧结果过期。ID查询仅读取近邻实际CSV；完全漏检事件可无ID录入并计作漏报。','backend/annotations.py','docs/visual-evaluation.md')
fig(7,'fig7-evaluation.png','人工评测结果区  受控流程验收，非真实精度实验',3.8)
p('事件按类型、映射ID和时间容差进行Hungarian一对一匹配，输出TP、FP、FN、precision和recall；默认容差0.15s可调整。无预测时precision、无真值时recall不可定义，未核验类别与速度保持unavailable。图7受控标签特意包含误报和未映射漏报，TP2/FP1/FN1只验证匹配契约，不代表真实交通报警性能。','backend/evaluation.py','docs/visual-evaluation.md','plans/002-browser-result.json','plans/002-verification.md')

newpage()
h('3 作品特色与创新',1)
p('以下特色属于当前系统设计与工程集成，不主张检测网络或基础跟踪算法创新。','backend/vision/detector.py','plans/002-verification.md')
h('3.1 观测与测量前提显式约束')
p('observed与标定区域门控避免将预测当测量，缺尺度则留空。回归测试及桥梁CSV的空世界坐标、空速度与零相关报警支持实现；明确车辆帧分母和窗口覆盖，使输出可解释，不构成精度证明。','backend/processor.py','backend/reports.py','tests/test_analysis_features.py','backend/data/runs/20261002-170236-c53990/vehicles.csv')
h('3.2 同帧展示与来源可追溯的复核')
p('图像与状态成对提交；截图、核验回放与人工记录分离，复核不调用实时定位。乱序解码测试及Range回放验收支持从二维状态返回原画面解释事件，截图不作为精度证据。','frontend/src/utils/snapshotState.ts','frontend/tests/snapshotState.test.mjs','backend/evidence.py','plans/2026-10-02-feature-verification.md')
h('3.3 机位复用与人工评测版本绑定')
p('模板默认清尺度，复制实测尺度须单独确认；标注绑定指纹与版本，修改使旧结果失效。比例边界、修订冲突、部分类别和无ID漏报测试及浏览器保存重开验收支持实现，迁移效果仍待验证。','backend/camera_templates.py','backend/annotations.py','tests/test_camera_templates.py','tests/test_annotations.py','plans/002-verification.md')
h('4 系统实现与实验验证',1)
h('4.1 实现环境 模型及演示素材')
p('本次环境为macOS 26.6 arm64、Python 3.13.0、FastAPI 0.136.3、PyTorch 2.11.0；前端为React 19、TypeScript、Vite、ECharts。已验证CPU，未实测NVIDIA GPU。主要素材为2160×3840、30FPS竖向桥梁录像，另有3840×2160横向素材验证方向与布局。','README.md','frontend/src/App.tsx','docs/review-smoke-2026-10-02.json','docs/twin-orientation.md')
h('4.2 功能流程与实际界面')
p('使用顺序为录像导入、道路配置与标定、分析及同帧展示、报告导出、独立回放和人工评测。图4、图6为真实桥梁界面，图7为受控验收。','plans/002-verification.md','plans/2026-10-02-feature-verification.md')
h('4.3 软件测试和运行验证')
p('2026年10月3日在隔离数据目录复测后端88项、前端27项均通过，并通过HTTP、可信来源空闲WebSocket和正常退出检查。表4区分本次复测、历史验收和实际输出。测试验证软件契约，不将其通过率换算成感知准确率。','tests/conftest.py','scripts/check_runtime.py','plans/002-verification.md')
table(4,'功能验证与实际运行证据',['验证层次','结果或材料','支持的结论及边界'],[
['本次软件回归','后端88通过；前端27通过','几何、规则、同步、版本与异常输入契约'],
['本次运行检查','HTTP / WebSocket / 正常退出通过','隔离合成素材；不证明交通精度'],
['历史实际浏览器','模板、报告、独立回放、人工评测','plans/002-verification.md；受控标签'],
['真实桥梁CSV','256帧；19次机器过线；28个ID','保存来源和配置；未经逐车真值核验'],
['观测样本口径','1074观测帧；1088预测帧','2162车辆帧行；有效观测速样本0'],
['既有CPU smoke','42.5FPS；相邻帧ID保留98.9%','单次离线吞吐和连续性；非精度']],[3.0,6.2,7.8],['plans/002-verification.md','scripts/check_runtime.py','docs/review-smoke-2026-10-02.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/vehicles.csv'])

newpage()
h('4.4 有真实真值支撑的评测状态')
p('尚无已保存的真实隧道或实测测速真值评测。运行段20261002-170236-c53990来源SHA256前24位与身份一致，配置完整、状态completed。图8绘制全部256行CSV：0.0–8.5s，累计过线19、最多在场13、确认轨迹ID共28个。曲线为机器输出，ID不等于物理车辆总数。','backend/data/runs/20261002-170236-c53990/run.json','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','backend/data/runs/20261002-170236-c53990/vehicles.csv','plans/002-verification.md')
fig(8,'fig8-traffic.png','真实桥梁运行CSV交通曲线  全部统计帧，未经人工真值核验',3.6)
p('该段1074条实际观测车辆帧中有效速度样本0条，世界坐标、速度与加速度为空，机器事件0条；这验证未标定门槛，不能推断场景不存在异常。末帧每分钟流率134.12辆/min由19÷8.5×60短时折算，不是完整一分钟计数。既有smoke的42.5FPS为256帧离线处理吞吐；30FPS是原视频帧率，浏览器显示帧率受推送、解码与渲染影响，当前未独立测量。98.9%相邻帧ID保留率仅表示运行连续性。','backend/data/runs/20261002-170236-c53990/vehicles.csv','backend/data/runs/20261002-170236-c53990/events.csv','backend/data/runs/20261002-170236-c53990/traffic_statistics.csv','docs/review-smoke-2026-10-02.json')
h('4.5 当前局限与后续验证计划')
p('现有桥梁录像不能覆盖隧道低照度、车灯反光、遮挡和相机震动；四点平面模型未补偿镜头畸变、道路高差及移动相机。需补充合法来源的固定隧道录像、机位与实测尺寸，独立完整标注各区间过线及事件，保存核验者、容差和版本，再报告计数绝对误差及事件TP/FP/FN。测速需独立参考测量和ID时间对应，报告匹配覆盖率、MAE、RMSE，不能仅报告匹配成功的样本。检测mAP和MOT类指标另需逐帧框及物理ID标注、相应导出；还需长时运行、GPU、干净安装及RTSP断流恢复验收。以上为后续方案，未作为已完成实验。','docs/accuracy-evaluation.md','docs/visual-evaluation.md','plans/2026-10-02-implementation-status.md')
h('5 应用前景与推广条件',1)
p('系统可用于监测教学、离线录像分析和事件人工复查。推广需取得固定机位与实测尺度，完成代表性隧道验证、现场阈值配置、来源授权、存储和故障处置。当前无落地案例、多摄像头并发或多用户权限，生产应用还需流接入、容量和持续运行保障。模型与视频许可须核实，不能承诺商业化直接使用。','models.json','backend/service.py','docs/camera-templates.md','plans/002-verification.md')
h('参考文献',1)
refs=[
('[1] Veres M, Moussa M. Deep Learning for Intelligent Transportation Systems: A Survey of Emerging Trends[J]. IEEE Transactions on Intelligent Transportation Systems, 2020, 21(8): 3152-3168. DOI: 10.1109/TITS.2019.2929020.','https://doi.org/10.1109/TITS.2019.2929020'),
('[2] Sochor J, Juránek R, Herout A. Traffic Surveillance Camera Calibration by 3D Model Bounding Box Alignment for Accurate Vehicle Speed Measurement[J]. Computer Vision and Image Understanding, 2017, 161: 87-98. DOI: 10.1016/j.cviu.2017.05.015.','https://doi.org/10.1016/j.cviu.2017.05.015'),
('[3] Sochor J, Juránek R, Špaňhel J, et al. Comprehensive Data Set for Automatic Single Camera Visual Speed Measurement[J]. IEEE Transactions on Intelligent Transportation Systems, 2019, 20(5): 1633-1643. DOI: 10.1109/TITS.2018.2825609.','https://doi.org/10.1109/TITS.2018.2825609'),
('[4] dronefreak. visdrone-yolo11n 模型卡[EB/OL]. [2026-10-03].','https://huggingface.co/dronefreak/visdrone-yolo11n'),
('[5] Ultralytics. Ultralytics YOLO11 官方文档[EB/OL]. [2026-10-03].','https://docs.ultralytics.com/models/yolo11/'),
('[6] Torchvision. ssdlite320_mobilenet_v3_large 官方文档[EB/OL]. [2026-10-03].','https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.detection.ssdlite320_mobilenet_v3_large.html'),
('[7] Zhu P, Wen L, Bian X, et al. Vision Meets Drones: A Challenge[EB/OL]. 2018. arXiv:1804.07437.','https://arxiv.org/abs/1804.07437'),
('[8] Bewley A, Ge Z, Ott L, et al. Simple Online and Realtime Tracking[C]. ICIP, 2016. DOI: 10.1109/ICIP.2016.7533003.','https://doi.org/10.1109/ICIP.2016.7533003'),
('[9] OpenCV. Geometric Image Transformations[EB/OL]. [2026-10-03].','https://docs.opencv.org/4.13.0/da/d54/group__imgproc__transform.html'),
]
from docx.opc.constants import RELATIONSHIP_TYPE as RT
for text,url in refs:
    pa=doc.add_paragraph(text);pa.paragraph_format.first_line_indent=Pt(0);pa.paragraph_format.space_after=Pt(1)
    for r in pa.runs:r.font.size=Pt(10.5)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz);rr.append(rp)
    tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=' 在线来源';rr.append(tt);link.append(rr);pa._p.append(link)


path=OUT/'基于视频感知的隧道交通数字孪生系统作品说明书.docx';doc.save(path)
(OUT/'evidence/manuscript.md').write_text('\n\n'.join(draft),encoding='utf-8')
with (OUT/'evidence/claims.csv').open('w',newline='',encoding='utf-8-sig') as f:
    wr=csv.DictWriter(f,fieldnames=list(claims[0]));wr.writeheader();wr.writerows(claims)
(OUT/'evidence/displays.json').write_text(json.dumps(displays,ensure_ascii=False,indent=2))
print(path)
